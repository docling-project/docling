# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for Word list numbering behaviour.

Kept separate from ``test_backend_msword.py`` so that file stays under the
repository's per-file line limit.
"""

from io import BytesIO

import pytest
from docling_core.types.doc import (
    DoclingDocument,
    DocumentOrigin,
    ListGroup,
    ListItem,
)
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docling.backend.msword_backend import (
    MsWordDocumentBackend,
    _format_enum_counter,
)
from docling.datamodel.base_models import InputFormat
from docling.datamodel.document import InputDocument


def test_ordered_list_resumes_numbering_after_intervening_list(tmp_path):
    """An ordered list interrupted by a bullet list must keep counting.

    Word numbers continuously per ``w:numId``, so items sharing one numId are a
    single list even when other content (including a list with a different
    numId) sits between them. Docling used to reset the counter whenever the
    numId changed, so the resumed items restarted at 1.
    """

    doc = Document()
    numbering = doc.part.numbering_part.element

    def add_numbering(abstract_id: str, num_id: str, num_fmt: str, lvl_text: str):
        abstract_num = OxmlElement("w:abstractNum")
        abstract_num.set(qn("w:abstractNumId"), abstract_id)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        lvl.append(fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), lvl_text)
        lvl.append(text)
        abstract_num.append(lvl)
        numbering.append(abstract_num)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), num_id)
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_id)
        num.append(ref)
        numbering.append(num)

    add_numbering("300", "301", "decimal", "%1.")
    add_numbering("400", "401", "bullet", "•")

    def add_item(text: str, num_id: str):
        paragraph = doc.add_paragraph(text, style="List Paragraph")
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_pr.append(ilvl)
        num_id_elem = OxmlElement("w:numId")
        num_id_elem.set(qn("w:val"), num_id)
        num_pr.append(num_id_elem)
        paragraph._element.get_or_add_pPr().append(num_pr)

    add_item("First ordered item", "301")
    add_item("Second ordered item", "301")
    add_item("bullet one", "401")
    add_item("bullet two", "401")
    add_item("Third ordered item", "301")

    docx_path = tmp_path / "resumed_ordered_list.docx"
    doc.save(str(docx_path))

    in_doc = InputDocument(
        path_or_stream=docx_path,
        format=InputFormat.DOCX,
        backend=MsWordDocumentBackend,
        filename=docx_path.name,
    )
    converted = MsWordDocumentBackend(in_doc=in_doc, path_or_stream=docx_path).convert()

    markers = [
        (item.text, item.marker)
        for item, _ in converted.iterate_items()
        if isinstance(item, ListItem)
    ]

    assert markers == [
        ("First ordered item", "1."),
        ("Second ordered item", "2."),
        ("bullet one", ""),
        ("bullet two", ""),
        ("Third ordered item", "3."),
    ]


def test_manage_list_structure_no_keyerror_when_use_level_exceeds_parents(tmp_path):
    """_manage_list_structure must not raise KeyError when use_level exceeds parents.

    The pathological state is:
      - parents has keys 0..11, with key 0 set to None (cleared) and others
        holding NodeItems from headings / earlier lists.
      - level_at_new_list is 11 (set by the previous list item of the same numId).
      - A second item for the same numId arrives with ilevel=2, which triggers
        the "New list sequence" branch and computes use_level = 11 + 2 = 13.
      - parents.get(12) is not a key at all → previously raised KeyError: 12.
    """

    docx_io = _make_empty_docx()
    docx_path = tmp_path / "empty.docx"
    docx_path.write_bytes(docx_io.getvalue())

    in_doc = InputDocument(
        path_or_stream=docx_path,
        format=InputFormat.DOCX,
        backend=MsWordDocumentBackend,
        filename=docx_path.name,
    )
    backend = MsWordDocumentBackend(in_doc=in_doc, path_or_stream=docx_path)

    out_doc = DoclingDocument(
        name="test",
        origin=DocumentOrigin(
            filename="test.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            binary_hash="0",
        ),
    )

    body_node = out_doc.body
    first_list_gr = out_doc.add_list_group(name="list", parent=body_node)
    backend.parents = dict.fromkeys(range(12), body_node)
    backend.parents[0] = None  # gap that makes _get_level() return 0
    backend.parents[11] = first_list_gr

    # Simulate history: the immediately previous item was numid=27, ilevel=2.
    backend.history = {
        "names": [None, "list"],
        "levels": [None, 11],
        "numids": [None, 27],
        "indents": [None, 2],
    }
    backend.level_at_new_list = 11
    backend.last_numid = 27

    # This must not raise KeyError.
    elem_ref, use_level = backend._manage_list_structure(
        doc=out_doc, numid=27, ilevel=2
    )

    assert isinstance(backend.parents.get(use_level), ListGroup)


def test_manage_list_structure_no_keyerror_open_indented_list_exceeds_parents(tmp_path):
    """_manage_list_structure must not raise KeyError in the "Open indented list" branch.

    The pathological state is:
      - parents has keys 0..11, with keys 0..10 holding body nodes and key 11
        holding a ListGroup.
      - level_at_new_list is 11, prev_indent is 2, ilevel is 4.
      - The "Open indented list" loop runs for i in range(14, 16), accessing
        parents[i - 1] where i - 1 = 13 is not a key in parents.
    """
    docx_path = tmp_path / "empty.docx"
    docx_path.write_bytes(_make_empty_docx().getvalue())

    in_doc = InputDocument(
        path_or_stream=docx_path,
        format=InputFormat.DOCX,
        backend=MsWordDocumentBackend,
        filename=docx_path.name,
    )
    backend = MsWordDocumentBackend(in_doc=in_doc, path_or_stream=docx_path)

    out_doc = DoclingDocument(
        name="test",
        origin=DocumentOrigin(
            filename="test.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            binary_hash="0",
        ),
    )

    body_node = out_doc.body
    first_list_gr = out_doc.add_list_group(name="list", parent=body_node)
    backend.parents = dict.fromkeys(range(12), body_node)
    backend.parents[11] = first_list_gr

    # Simulate history: same numId, previous item was at ilevel=2.
    backend.history = {
        "names": [None, "list"],
        "levels": [None, 11],
        "numids": [None, 27],
        "indents": [None, 2],
    }
    backend.level_at_new_list = 11
    backend.last_numid = 27

    elem_ref, use_level = backend._manage_list_structure(
        doc=out_doc, numid=27, ilevel=4
    )

    assert isinstance(backend.parents.get(use_level), ListGroup)


def test_list_markers_follow_num_fmt_end_to_end(tmp_path):
    """Lettered/roman markers must survive convert(), including plain-suffix lvlText.

    Word keeps the counter as an integer and the display form in ``w:numFmt``.
    A level whose ``lvlText`` is only ``%2)`` (no extra text after stripping
    placeholders/punctuation) must still take the lvlText path when numFmt is
    non-decimal, so ``lowerLetter`` renders ``a)`` rather than hierarchical
    ``1.a.``. Hierarchical decimal/letter/roman mixes stay on the fallback form.
    """

    doc = Document()
    numbering = doc.part.numbering_part.element

    def add_numbering(abstract_id: str, num_id: str, levels: list[tuple[str, str]]):
        abstract_num = OxmlElement("w:abstractNum")
        abstract_num.set(qn("w:abstractNumId"), abstract_id)
        for ilvl, (num_fmt, lvl_text) in enumerate(levels):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(ilvl))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            lvl.append(start)
            fmt = OxmlElement("w:numFmt")
            fmt.set(qn("w:val"), num_fmt)
            lvl.append(fmt)
            text_el = OxmlElement("w:lvlText")
            text_el.set(qn("w:val"), lvl_text)
            lvl.append(text_el)
            abstract_num.append(lvl)
        numbering.append(abstract_num)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), num_id)
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_id)
        num.append(ref)
        numbering.append(num)

    # Hierarchical mix: decimal / lowerLetter / upperRoman (fallback form).
    add_numbering(
        "500",
        "501",
        [("decimal", "%1."), ("lowerLetter", "%1.%2."), ("upperRoman", "%1.%2.%3.")],
    )
    # Plain-suffix lowerLetter: the issue's a) / b) case (lvlText "%1)").
    add_numbering("600", "601", [("lowerLetter", "%1)")])
    # Other non-decimal formats with punctuation-only templates.
    add_numbering("700", "701", [("upperLetter", "%1.")])
    add_numbering("800", "801", [("lowerRoman", "%1.")])
    add_numbering("900", "901", [("upperRoman", "%1.")])
    add_numbering("1000", "1001", [("decimalZero", "%1.")])

    def add_item(text: str, num_id: str, ilvl_val: int = 0):
        paragraph = doc.add_paragraph(text, style="List Paragraph")
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(ilvl_val))
        num_pr.append(ilvl)
        num_id_elem = OxmlElement("w:numId")
        num_id_elem.set(qn("w:val"), num_id)
        num_pr.append(num_id_elem)
        paragraph._element.get_or_add_pPr().append(num_pr)

    add_item("top one", "501", 0)
    add_item("lettered first", "501", 1)
    add_item("lettered second", "501", 1)
    add_item("roman first", "501", 2)
    add_item("roman second", "501", 2)
    add_item("top two", "501", 0)

    add_item("plain suffix first", "601")
    add_item("plain suffix second", "601")

    add_item("upper letter", "701")
    add_item("lower roman", "801")
    add_item("upper roman", "901")
    add_item("zero pad", "1001")

    docx_path = tmp_path / "num_fmt_markers.docx"
    doc.save(str(docx_path))

    in_doc = InputDocument(
        path_or_stream=docx_path,
        format=InputFormat.DOCX,
        backend=MsWordDocumentBackend,
        filename=docx_path.name,
    )
    converted = MsWordDocumentBackend(in_doc=in_doc, path_or_stream=docx_path).convert()

    markers = [
        (item.text, item.marker)
        for item, _ in converted.iterate_items()
        if isinstance(item, ListItem)
    ]

    assert markers == [
        ("top one", "1."),
        ("lettered first", "1.a."),
        ("lettered second", "1.b."),
        ("roman first", "1.b.I."),
        ("roman second", "1.b.II."),
        ("top two", "2."),
        ("plain suffix first", "a)"),
        ("plain suffix second", "b)"),
        ("upper letter", "A."),
        ("lower roman", "i."),
        ("upper roman", "I."),
        ("zero pad", "01."),
    ]


# Expected strings are the markers Microsoft Word 16.112 (macOS) renders for a
# level with the given ``w:numFmt`` and ``w:start``. Character sets are from
# ECMA-376-1:2016 §17.18.59; the rule each format follows is noted inline.
# Zero glyphs look alike, so they are written as escapes: U+25CB for
# chineseCounting, U+3007 for chineseCountingThousand and ideographDigital,
# U+96F6 (零) for chineseLegalSimplified.
_BOUNDARY_COUNTERS = (
    1,
    9,
    10,
    11,
    19,
    20,
    21,
    99,
    100,
    101,
    110,
    999,
    1000,
    1001,
    10000,
    10005,
    100000,
)
_EAST_ASIAN_MARKERS: dict[str, tuple[str, ...]] = {
    # 十 reading to 99, digit by digit with U+25CB from 100 (ECMA-376 example
    # 九十九, 一○○, 一○一; Word identical, no upper limit).
    "chineseCounting": (
        "一",
        "九",
        "十",
        "十一",
        "十九",
        "二十",
        "二十一",
        "九十九",
        "一\u25cb\u25cb",
        "一\u25cb一",
        "一一\u25cb",
        "九九九",
        "一\u25cb\u25cb\u25cb",
        "一\u25cb\u25cb一",
        "一\u25cb\u25cb\u25cb\u25cb",
        "一\u25cb\u25cb\u25cb五",
        "一\u25cb\u25cb\u25cb\u25cb\u25cb",
    ),
    # Place-value reading; 一 dropped only for a bare 10-19; U+3007 marks every
    # zero run followed by a digit, also right after 万 (Word: 一万〇五). The
    # standard's example 一十 and its zero 零 are not what Word renders.
    "chineseCountingThousand": (
        "一",
        "九",
        "十",
        "十一",
        "十九",
        "二十",
        "二十一",
        "九十九",
        "一百",
        "一百\u3007一",
        "一百一十",
        "九百九十九",
        "一千",
        "一千\u3007一",
        "一万",
        "一万\u3007五",
        "一十万",
    ),
    # Banker's digits with 拾佰仟, coefficient never dropped (10 = 壹拾), zero
    # 零, ten thousand U+842C as in Word ([MS-OI29500] 2.1.548 q).
    "chineseLegalSimplified": (
        "壹",
        "玖",
        "壹拾",
        "壹拾壹",
        "壹拾玖",
        "贰拾",
        "贰拾壹",
        "玖拾玖",
        "壹佰",
        "壹佰零壹",
        "壹佰壹拾",
        "玖佰玖拾玖",
        "壹仟",
        "壹仟零壹",
        "壹萬",
        "壹萬零伍",
        "壹拾萬",
    ),
    # No zero inside a number; 一 omitted before 十, 百 and a leading 千, kept
    # as 一万 and as 一千 after a 万 group (Word renders; ECMA-376 shows only
    # the pattern to 二十一).
    "japaneseCounting": (
        "一",
        "九",
        "十",
        "十一",
        "十九",
        "二十",
        "二十一",
        "九十九",
        "百",
        "百一",
        "百十",
        "九百九十九",
        "千",
        "千一",
        "一万",
        "一万五",
        "十万",
    ),
    # Digit substitution with U+3007 (ECMA-376; Word identical).
    "ideographDigital": (
        "一",
        "九",
        "一\u3007",
        "一一",
        "一九",
        "二\u3007",
        "二一",
        "九九",
        "一\u3007\u3007",
        "一\u3007一",
        "一一\u3007",
        "九九九",
        "一\u3007\u3007\u3007",
        "一\u3007\u3007一",
        "一\u3007\u3007\u3007\u3007",
        "一\u3007\u3007\u3007五",
        "一\u3007\u3007\u3007\u3007\u3007",
    ),
    # Fullwidth digits (ECMA-376; Word identical).
    "decimalFullWidth": tuple(
        "".join(chr(0xFF10 + int(char)) for char in str(counter))
        for counter in _BOUNDARY_COUNTERS
    ),
    # Ten Heavenly Stems, then the decimal string (ECMA-376 example 癸, 11, 12;
    # Word identical).
    "ideographTraditional": (
        "甲",
        "壬",
        "癸",
        "11",
        "19",
        "20",
        "21",
        "99",
        "100",
        "101",
        "110",
        "999",
        "1000",
        "1001",
        "10000",
        "10005",
        "100000",
    ),
    # Twelve Earthly Branches, then the decimal string (ECMA-376 example 亥,
    # 13, 14; Word identical for the fallback).
    "ideographZodiac": (
        "子",
        "申",
        "酉",
        "戌",
        "19",
        "20",
        "21",
        "99",
        "100",
        "101",
        "110",
        "999",
        "1000",
        "1001",
        "10000",
        "10005",
        "100000",
    ),
    # U+2460-U+2473 for 1-20, then the decimal string (ECMA-376; Word identical).
    "decimalEnclosedCircle": (
        "①",
        "⑨",
        "⑩",
        "⑪",
        "⑲",
        "⑳",
        "21",
        "99",
        "100",
        "101",
        "110",
        "999",
        "1000",
        "1001",
        "10000",
        "10005",
        "100000",
    ),
}


@pytest.mark.parametrize(
    ("num_fmt", "counter", "expected"),
    [
        (num_fmt, counter, expected)
        for num_fmt, markers in _EAST_ASIAN_MARKERS.items()
        for counter, expected in zip(_BOUNDARY_COUNTERS, markers, strict=True)
    ]
    + [
        # Zero is a valid w:start value.
        ("chineseCounting", 0, "\u25cb"),
        ("chineseCountingThousand", 0, "\u3007"),
        ("chineseLegalSimplified", 0, "零"),
        ("japaneseCounting", 0, "\u3007"),
        ("ideographDigital", 0, "\u3007"),
        ("decimalFullWidth", 0, "\uff10"),
        ("ideographTraditional", 0, "0"),
        ("decimalEnclosedCircle", 0, "0"),
        # Zero runs inside and across the 万 group.
        ("chineseCountingThousand", 1010, "一千\u3007一十"),
        ("chineseCountingThousand", 10101, "一万\u3007一百\u3007一"),
        ("chineseCountingThousand", 100010, "一十万\u3007一十"),
        ("chineseCountingThousand", 909090, "九十万\u3007九千\u3007九十"),
        ("chineseLegalSimplified", 100100, "壹拾萬零壹佰"),
        ("chineseLegalSimplified", 909090, "玖拾萬零玖仟零玖拾"),
        ("japaneseCounting", 1100, "千百"),
        ("japaneseCounting", 11100, "一万一千百"),
        ("japaneseCounting", 111111, "十一万一千百十一"),
        # Word renders an empty marker from 1,000,000 ([MS-OI29500] 2.1.548 j).
        ("chineseCountingThousand", 999999, "九十九万九千九百九十九"),
        ("chineseCountingThousand", 1000000, ""),
        ("chineseLegalSimplified", 1000000, ""),
        ("japaneseCounting", 1000000, ""),
        ("chineseCounting", 1000000, "一\u25cb\u25cb\u25cb\u25cb\u25cb\u25cb"),
    ],
)
def test_format_enum_counter_east_asian_num_fmt(num_fmt, counter, expected):
    assert _format_enum_counter(counter, num_fmt) == expected


def test_list_markers_follow_east_asian_num_fmt_end_to_end(tmp_path):
    """East Asian ``w:numFmt`` levels must keep their markers through convert().

    Chinese regulations and contracts number articles with ``第%1条`` and a
    ``%2`` wrapped in fullwidth parentheses. These formats used to be treated as non-numbered, so the list
    items came out as bullets with an empty marker and the article numbers were
    lost.
    """

    doc = Document()
    numbering = doc.part.numbering_part.element

    def add_numbering(
        abstract_id: str, num_id: str, levels: list[tuple[str, str, int]]
    ):
        abstract_num = OxmlElement("w:abstractNum")
        abstract_num.set(qn("w:abstractNumId"), abstract_id)
        for ilvl, (num_fmt, lvl_text, start_val) in enumerate(levels):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(ilvl))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), str(start_val))
            lvl.append(start)
            fmt = OxmlElement("w:numFmt")
            fmt.set(qn("w:val"), num_fmt)
            lvl.append(fmt)
            text_el = OxmlElement("w:lvlText")
            text_el.set(qn("w:val"), lvl_text)
            lvl.append(text_el)
            abstract_num.append(lvl)
        numbering.append(abstract_num)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), num_id)
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_id)
        num.append(ref)
        numbering.append(num)

    # Articles start at 9 so the list crosses the 十 boundary.
    add_numbering(
        "500",
        "501",
        [
            ("chineseCountingThousand", "第%1条", 9),
            ("chineseCounting", "\uff08%2\uff09", 1),
            ("decimalEnclosedCircle", "%3", 1),
        ],
    )
    # Punctuation-only templates must also take the lvlText path.
    add_numbering("600", "601", [("chineseLegalSimplified", "%1、", 1)])
    add_numbering("700", "701", [("ideographTraditional", "%1.", 1)])
    add_numbering("800", "801", [("japaneseCounting", "%1", 100)])

    def add_item(text: str, num_id: str, ilvl_val: int = 0):
        paragraph = doc.add_paragraph(text, style="List Paragraph")
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(ilvl_val))
        num_pr.append(ilvl)
        num_id_elem = OxmlElement("w:numId")
        num_id_elem.set(qn("w:val"), num_id)
        num_pr.append(num_id_elem)
        paragraph._element.get_or_add_pPr().append(num_pr)

    add_item("article nine", "501", 0)
    add_item("first clause", "501", 1)
    add_item("second clause", "501", 1)
    add_item("first point", "501", 2)
    add_item("article ten", "501", 0)
    add_item("article eleven", "501", 0)

    add_item("legal one", "601")
    add_item("legal two", "601")
    add_item("stem one", "701")
    add_item("japanese hundred", "801")

    docx_path = tmp_path / "east_asian_num_fmt_markers.docx"
    doc.save(str(docx_path))

    in_doc = InputDocument(
        path_or_stream=docx_path,
        format=InputFormat.DOCX,
        backend=MsWordDocumentBackend,
        filename=docx_path.name,
    )
    converted = MsWordDocumentBackend(in_doc=in_doc, path_or_stream=docx_path).convert()

    items = [
        (item.text, item.marker, item.enumerated)
        for item, _ in converted.iterate_items()
        if isinstance(item, ListItem)
    ]

    assert items == [
        ("article nine", "第九条", True),
        ("first clause", "\uff08一\uff09", True),
        ("second clause", "\uff08二\uff09", True),
        ("first point", "①", True),
        ("article ten", "第十条", True),
        ("article eleven", "第十一条", True),
        ("legal one", "壹、", True),
        ("legal two", "贰、", True),
        ("stem one", "甲.", True),
        ("japanese hundred", "百", True),
    ]


def _make_empty_docx():
    """Return an in-memory .docx with no content."""

    buf = BytesIO()
    Document().save(buf)
    buf.seek(0)
    return buf
