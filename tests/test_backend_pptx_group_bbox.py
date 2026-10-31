# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for the provenance boxes of shapes inside PowerPoint groups."""

from pathlib import Path

from docling_core.types.doc import TextItem
from pptx import Presentation
from pptx.util import Emu, Inches

from .test_backend_pptx import get_converter


def _boxes_by_text(path: Path) -> dict[str, tuple[float, float, float, float]]:
    doc = get_converter().convert(path).document
    return {
        item.text: (
            item.prov[0].bbox.l,
            item.prov[0].bbox.t,
            item.prov[0].bbox.r,
            item.prov[0].bbox.b,
        )
        for item, _ in doc.iterate_items()
        if isinstance(item, TextItem)
    }


def _add_text(shapes, text: str, left, top, width, height):
    box = shapes.add_textbox(left, top, width, height)
    box.text_frame.text = text
    return box


def test_pptx_grouped_shape_bbox_follows_group_transform(tmp_path: Path):
    """A shape in a moved, resized or flipped group gets its drawn position.

    python-pptx reports grouped shapes in the group's child coordinates
    (``a:chOff``/``a:chExt``); the provenance box must be mapped through the
    group's ``a:off``/``a:ext`` (and flips) of every enclosing group.
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text(slide.shapes, "plain", Inches(0.5), Inches(0.5), Inches(2), Inches(1))

    moved = slide.shapes.add_group_shape()
    _add_text(moved.shapes, "moved", Inches(1), Inches(2), Inches(1), Inches(1))
    moved.left, moved.top = Emu(Inches(4)), Emu(Inches(3))

    resized = slide.shapes.add_group_shape()
    _add_text(resized.shapes, "resized-a", Inches(1), Inches(5), Inches(1), Inches(1))
    _add_text(resized.shapes, "resized-b", Inches(2), Inches(5), Inches(1), Inches(1))
    resized.width = Emu(Inches(4))

    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    _add_text(inner.shapes, "nested", Inches(6), Inches(1), Inches(1), Inches(1))
    _add_text(outer.shapes, "outer", Inches(7), Inches(1), Inches(1), Inches(1))
    inner.left = Emu(Inches(6.5))
    outer.top = Emu(Inches(2))
    outer.height = Emu(Inches(2))

    flipped = slide.shapes.add_group_shape()
    _add_text(flipped.shapes, "flip-left", Inches(1), Inches(6), Inches(1), Inches(1))
    _add_text(flipped.shapes, "flip-right", Inches(3), Inches(6), Inches(1), Inches(1))
    flipped._element.grpSpPr.get_or_add_xfrm().set("flipH", "1")

    path = tmp_path / "groups.pptx"
    prs.save(path)
    boxes = _boxes_by_text(path)

    def inches(left, top, right, bottom):
        return tuple(Inches(value) for value in (left, top, right, bottom))

    assert boxes == {
        "plain": inches(0.5, 0.5, 2.5, 1.5),
        "moved": inches(4, 3, 5, 4),
        "resized-a": inches(1, 5, 3, 6),
        "resized-b": inches(3, 5, 5, 6),
        # The inner group moves right by 0.5 in, then the outer group starts
        # at 2 in and doubles in height.
        "nested": inches(6.5, 2, 7.5, 4),
        "outer": inches(7, 2, 8, 4),
        "flip-left": inches(3, 6, 4, 7),
        "flip-right": inches(1, 6, 2, 7),
    }
