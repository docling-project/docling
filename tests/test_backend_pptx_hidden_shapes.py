# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for shapes hidden in PowerPoint's Selection Pane."""

from pathlib import Path

import pytest
from docling_core.types.doc import ContentLayer, TextItem

from docling.datamodel.document import DoclingDocument

from .test_backend_pptx import get_converter


def _set_shape_hidden(shape, value: str = "1") -> None:
    """Hide a shape the way PowerPoint's Selection Pane does."""
    shape._element.xpath("./*/p:cNvPr")[0].set("hidden", value)


def _layers_by_text(doc: DoclingDocument) -> dict[str, str]:
    return {
        item.text: item.content_layer.value
        for item, _ in doc.iterate_items(included_content_layers=set(ContentLayer))
        if isinstance(item, TextItem)
    }


@pytest.mark.parametrize("hidden_value", ["1", "true"])
def test_pptx_hidden_shape_goes_to_invisible_layer(tmp_path: Path, hidden_value: str):
    """A shape hidden in the Selection Pane is left out of the default export.

    Its text stays available in ``ContentLayer.INVISIBLE``, the same treatment
    a hidden slide gets. ``hidden="0"`` and a missing attribute stay visible.
    """
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shown = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    shown.text_frame.text = "Shown shape text"
    hidden = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    hidden.text_frame.text = "Hidden shape text"
    _set_shape_hidden(hidden, hidden_value)
    not_hidden = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
    not_hidden.text_frame.text = "Explicitly shown text"
    _set_shape_hidden(not_hidden, "0")

    pptx_path = tmp_path / "hidden_shape.pptx"
    prs.save(pptx_path)
    doc = get_converter().convert(pptx_path).document

    assert _layers_by_text(doc) == {
        "Shown shape text": "body",
        "Hidden shape text": "invisible",
        "Explicitly shown text": "body",
    }
    assert doc.export_to_markdown() == "Shown shape text\n\nExplicitly shown text"


def test_pptx_hidden_group_hides_everything_inside_it(tmp_path: Path):
    """Hiding a group hides its members; hiding one member leaves its siblings."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    hidden_group = slide.shapes.add_group_shape()
    for index, text in enumerate(["Hidden group first", "Hidden group second"]):
        box = hidden_group.shapes.add_textbox(
            Inches(1), Inches(1 + index), Inches(4), Inches(0.5)
        )
        box.text_frame.text = text
    _set_shape_hidden(hidden_group)

    shown_group = slide.shapes.add_group_shape()
    kept = shown_group.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(0.5))
    kept.text_frame.text = "Shown group kept"
    dropped = shown_group.shapes.add_textbox(
        Inches(1), Inches(5), Inches(4), Inches(0.5)
    )
    dropped.text_frame.text = "Shown group hidden member"
    _set_shape_hidden(dropped)

    pptx_path = tmp_path / "hidden_group.pptx"
    prs.save(pptx_path)
    doc = get_converter().convert(pptx_path).document

    assert _layers_by_text(doc) == {
        "Hidden group first": "invisible",
        "Hidden group second": "invisible",
        "Shown group kept": "body",
        "Shown group hidden member": "invisible",
    }
    assert doc.export_to_markdown() == "Shown group kept"


def test_pptx_hidden_table_and_picture_go_to_invisible_layer(tmp_path: Path):
    """The hidden flag applies to graphic frames and pictures, not only text."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    image_path = tmp_path / "pixel.png"
    Image.new("RGB", (40, 20), "red").save(image_path)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    frame.table.cell(0, 0).text = "Hidden cell"
    frame.table.cell(0, 1).text = "Other cell"
    _set_shape_hidden(frame)
    picture = slide.shapes.add_picture(str(image_path), Inches(1), Inches(3))
    _set_shape_hidden(picture)
    shown = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    shown.text_frame.text = "Shown text"

    pptx_path = tmp_path / "hidden_table_picture.pptx"
    prs.save(pptx_path)
    doc = get_converter().convert(pptx_path).document

    assert len(doc.tables) == 1
    assert len(doc.pictures) == 1
    assert doc.tables[0].content_layer == ContentLayer.INVISIBLE
    assert doc.pictures[0].content_layer == ContentLayer.INVISIBLE
    assert doc.export_to_markdown() == "Shown text"
