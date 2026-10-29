# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for PowerPoint shapes filled with a picture (``a:blipFill``)."""

import warnings
from io import BytesIO
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from .test_backend_pptx import get_converter


def _picture_filled_shape_pptx(
    path: Path, text: str = "", embed: bool = True, r_id: str | None = None
):
    """Save a one-slide PPTX with a rectangle filled with a 64x32 PNG.

    ``r_id`` overrides the relationship id the fill points to.
    """
    png = BytesIO()
    Image.new("RGB", (64, 32), "red").save(png, "PNG")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
    )
    if text:
        shape.text_frame.text = text
    _, image_r_id = shape.part.get_or_add_image_part(BytesIO(png.getvalue()))

    blip_fill = etree.Element(qn("a:blipFill"))
    blip = etree.SubElement(blip_fill, qn("a:blip"))
    blip.set(qn("r:embed" if embed else "r:link"), r_id or image_r_id)
    etree.SubElement(etree.SubElement(blip_fill, qn("a:stretch")), qn("a:fillRect"))
    shape._element.spPr.find(qn("a:prstGeom")).addnext(blip_fill)
    prs.save(path)


def test_pptx_picture_filled_shape_is_extracted(tmp_path: Path):
    """A shape filled with a picture gives a picture item."""
    pptx_path = tmp_path / "blipfill.pptx"
    _picture_filled_shape_pptx(pptx_path)

    doc = get_converter().convert(pptx_path).document

    assert len(doc.pictures) == 1
    picture = doc.pictures[0]
    assert picture.image is not None
    assert picture.image.mimetype == "image/png"
    assert picture.image.size.width == 64.0
    assert picture.image.size.height == 32.0
    assert picture.prov and picture.prov[0].page_no == 1
    assert picture.prov[0].bbox.width > 0


def test_pptx_picture_filled_shape_keeps_its_text(tmp_path: Path):
    """The text of a picture-filled shape is kept next to the picture."""
    pptx_path = tmp_path / "blipfill_text.pptx"
    _picture_filled_shape_pptx(pptx_path, text="Caption on the fill")

    doc = get_converter().convert(pptx_path).document

    assert len(doc.pictures) == 1
    assert [t.text for t in doc.texts] == ["Caption on the fill"]


def test_pptx_linked_picture_fill_is_ignored(tmp_path: Path):
    """A picture fill that only links an external image adds no picture."""
    pptx_path = tmp_path / "blipfill_link.pptx"
    _picture_filled_shape_pptx(pptx_path, text="Linked fill", embed=False)

    with pytest.warns(UserWarning, match="picture fill without an embedded image"):
        doc = get_converter().convert(pptx_path).document

    assert len(doc.pictures) == 0
    assert [t.text for t in doc.texts] == ["Linked fill"]


def test_pptx_picture_fill_with_missing_part_is_skipped(tmp_path: Path):
    """A fill pointing at a missing relationship is skipped, text is kept."""
    pptx_path = tmp_path / "blipfill_dangling.pptx"
    _picture_filled_shape_pptx(pptx_path, text="Dangling fill", r_id="rId999")

    with pytest.warns(UserWarning, match="Skipping malformed picture shape"):
        doc = get_converter().convert(pptx_path).document

    assert len(doc.pictures) == 0
    assert [t.text for t in doc.texts] == ["Dangling fill"]


@pytest.mark.parametrize(
    "xml",
    [
        # No p:spPr at all.
        '<p:sp xmlns:p="{p}"/>',
        # p:spPr with a non-picture fill.
        '<p:sp xmlns:p="{p}" xmlns:a="{a}"><p:spPr><a:solidFill/></p:spPr></p:sp>',
    ],
    ids=["no_sppr", "solid_fill"],
)
def test_picture_fill_rid_without_picture_fill(xml: str):
    """Shapes without a picture fill give no relationship id and no warning."""
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        assert _picture_fill_rid_from_xml(xml) is None


def test_picture_fill_rid_with_empty_embed_warns():
    """A picture fill with an empty r:embed gives no relationship id."""
    xml = (
        '<p:sp xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}"><p:spPr><a:blipFill>'
        '<a:blip r:embed=""/></a:blipFill></p:spPr></p:sp>'
    )
    with pytest.warns(UserWarning, match="picture fill without an embedded image"):
        assert _picture_fill_rid_from_xml(xml) is None


def _picture_fill_rid_from_xml(xml: str):
    from types import SimpleNamespace

    from docling.backend.mspowerpoint_backend import MsPowerpointDocumentBackend

    element = etree.fromstring(
        xml.format(
            p="http://schemas.openxmlformats.org/presentationml/2006/main",
            a="http://schemas.openxmlformats.org/drawingml/2006/main",
            r="http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        )
    )
    return MsPowerpointDocumentBackend._picture_fill_rid(
        SimpleNamespace(_element=element)
    )
