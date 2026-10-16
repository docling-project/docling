# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from docling.backend.mspowerpoint_backend import MsPowerpointDocumentBackend
from docling.datamodel.base_models import InputFormat
from docling.document_converter import DocumentConverter


def get_converter() -> DocumentConverter:
    return DocumentConverter(allowed_formats=[InputFormat.PPTX])


def test_pptx_preserves_run_level_formatting_and_hyperlinks(tmp_path: Path):
    """Each run in a paragraph keeps its own formatting and hyperlink."""

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(4), Inches(1))
    p = box.text_frame.paragraphs[0]

    bold_run = p.add_run()
    bold_run.text = "Bold"
    bold_run.font.bold = True

    plain_run = p.add_run()
    plain_run.text = " plain "

    link_run = p.add_run()
    link_run.text = "link"
    link_run.hyperlink.address = "https://example.com/docs"

    pptx_path = tmp_path / "mixed_runs.pptx"
    prs.save(pptx_path)

    doc = get_converter().convert(pptx_path).document

    texts = [t for t in doc.texts if t.text]
    assert [t.text for t in texts] == ["Bold", " plain ", "link"]

    assert texts[0].formatting is not None
    assert texts[0].formatting.bold is True
    assert texts[0].hyperlink is None

    assert texts[1].formatting is not None
    assert texts[1].formatting.bold is False
    assert texts[1].hyperlink is None

    assert texts[2].formatting is not None
    assert texts[2].formatting.bold is False
    assert texts[2].hyperlink is not None
    assert str(texts[2].hyperlink) == "https://example.com/docs"


@pytest.mark.parametrize("malformation", ["dangling", "empty_rid", "no_rid"])
def test_pptx_malformed_run_hyperlink_is_dropped(tmp_path: Path, malformation: str):
    """A run whose hyperlink relationship cannot be resolved keeps its text."""

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(4), Inches(1))
    p = box.text_frame.paragraphs[0]

    broken_run = p.add_run()
    broken_run.text = "broken"
    broken_run.font.bold = True
    broken_run.hyperlink.address = "https://example.com/broken"

    p.add_run().text = " and "

    valid_run = p.add_run()
    valid_run.text = "valid"
    valid_run.hyperlink.address = "https://example.com/valid"

    hlink_click = broken_run._r.rPr.hlinkClick
    if malformation == "dangling":
        slide.part.rels.pop(hlink_click.rId)
    elif malformation == "empty_rid":
        hlink_click.set(qn("r:id"), "")
    else:
        del hlink_click.attrib[qn("r:id")]

    pptx_path = tmp_path / "malformed_hyperlink.pptx"
    prs.save(pptx_path)

    doc = get_converter().convert(pptx_path).document

    texts = [t for t in doc.texts if t.text]
    assert [t.text for t in texts] == ["broken", " and ", "valid"]
    assert texts[0].formatting is not None
    assert texts[0].formatting.bold is True
    assert texts[0].hyperlink is None
    assert texts[1].hyperlink is None
    assert str(texts[2].hyperlink) == "https://example.com/valid"


def test_pptx_paragraph_runs_are_read_once(tmp_path: Path, monkeypatch):
    """Each paragraph is split into runs once, including in whitespace-only shapes."""

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for text in ["first", "   "]:
        box = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(4), Inches(1))
        box.text_frame.text = text
    pptx_path = tmp_path / "paragraph_runs.pptx"
    prs.save(pptx_path)

    calls: list[str] = []
    iter_runs = MsPowerpointDocumentBackend._iter_paragraph_runs

    def spy(self, paragraph):
        calls.append(paragraph.text)
        return iter_runs(self, paragraph)

    monkeypatch.setattr(MsPowerpointDocumentBackend, "_iter_paragraph_runs", spy)

    get_converter().convert(pptx_path)

    assert sorted(calls) == ["   ", "first"]
