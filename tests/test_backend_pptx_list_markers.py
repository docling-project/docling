# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for the markers of auto-numbered PowerPoint list items."""

from pathlib import Path

import pytest
from docling_core.types.doc import ListItem
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from .test_backend_pptx import get_converter


def _save_numbered_deck(
    pptx_path: Path, lists: list[tuple[dict[str, str], list[str]]]
) -> None:
    """Save one slide with a text box per list, numbered with ``a:buAutoNum``."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for top, (attrs, texts) in enumerate(lists, start=1):
        box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(4), Inches(1))
        box.text_frame.text = texts[0]
        for text in texts[1:]:
            box.text_frame.add_paragraph().text = text
        for paragraph in box.text_frame.paragraphs:
            paragraph._p.get_or_add_pPr().append(
                paragraph._p.makeelement(qn("a:buAutoNum"), attrs)
            )
    prs.save(pptx_path)


def test_pptx_numbered_list_marker_follows_auto_number_type(tmp_path: Path):
    """A numbered item gets the marker of its ``a:buAutoNum/@type`` scheme.

    PowerPoint shows these lists as "a. b. c.", "I. II." or "(C) (D)", but every
    item used to get a decimal number and a period. Arabic schemes such as
    ``arabicParenR`` keep the decimal marker, as numbered DOCX lists do.
    """
    pptx_path = tmp_path / "auto_number_types.pptx"
    _save_numbered_deck(
        pptx_path,
        [
            ({"type": "alphaLcPeriod"}, ["Plan", "Build", "Test"]),
            ({"type": "romanUcPeriod"}, ["Scope", "Risks", "Costs", "Staff"]),
            ({"type": "arabicParenR"}, ["Draft", "Review"]),
            ({"type": "alphaUcParenBoth", "startAt": "3"}, ["Annex", "Index"]),
        ],
    )

    doc = get_converter().convert(pptx_path).document

    assert [
        (item.marker, item.text) for item in doc.texts if isinstance(item, ListItem)
    ] == [
        ("a.", "Plan"),
        ("b.", "Build"),
        ("c.", "Test"),
        ("I.", "Scope"),
        ("II.", "Risks"),
        ("III.", "Costs"),
        ("IV.", "Staff"),
        ("1.", "Draft"),
        ("2.", "Review"),
        ("(C)", "Annex"),
        ("(D)", "Index"),
    ]
    assert doc.export_to_markdown() == (
        "- a. Plan\n- b. Build\n- c. Test\n\n"
        "- I. Scope\n- II. Risks\n- III. Costs\n- IV. Staff\n\n"
        "1. Draft\n2. Review\n\n"
        "- (C) Annex\n- (D) Index"
    )


@pytest.mark.parametrize("start_at", ["0", "32768", "2000000000"])
def test_pptx_numbered_list_ignores_out_of_range_start_at(
    tmp_path: Path, start_at: str
):
    """A ``startAt`` outside 1 to 32767 is invalid, so the list starts at 1.

    Used as is, 2000000000 would build letter markers of about 77 million
    characters.
    """
    pptx_path = tmp_path / "out_of_range_start_at.pptx"
    _save_numbered_deck(
        pptx_path,
        [({"type": "alphaLcPeriod", "startAt": start_at}, ["Plan", "Build"])],
    )

    doc = get_converter().convert(pptx_path).document

    assert [
        (item.marker, item.text) for item in doc.texts if isinstance(item, ListItem)
    ] == [("a.", "Plan"), ("b.", "Build")]


def test_pptx_numbered_list_marker_follows_layout_auto_number_type(tmp_path: Path):
    """A scheme that the slide layout sets applies like one on the paragraph."""
    prs = Presentation()
    layout = prs.slide_layouts[1]
    layout_style = layout.placeholders[1]._element.find(".//" + qn("a:lstStyle"))
    assert layout_style is not None
    level = layout_style.makeelement(qn("a:lvl1pPr"))
    level.append(level.makeelement(qn("a:buAutoNum"), {"type": "romanLcParenR"}))
    layout_style.append(level)
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Steps"
    body = slide.placeholders[1].text_frame
    body.text = "Prepare"
    body.add_paragraph().text = "Ship"
    pptx_path = tmp_path / "layout_auto_number.pptx"
    prs.save(pptx_path)

    doc = get_converter().convert(pptx_path).document

    assert doc.export_to_markdown() == "# Steps\n\n- i) Prepare\n- ii) Ship"
