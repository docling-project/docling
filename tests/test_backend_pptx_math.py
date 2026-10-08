# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

from io import BytesIO

import pytest
from docling_core.types.doc import DocItemLabel, DoclingDocument, ListItem
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from docling.datamodel.base_models import DocumentStream, InputFormat
from docling.document_converter import DocumentConverter

_NAMESPACES = (
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" '
    'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
)
_EQUATION = "<a14:m><m:oMath><m:r><m:t>E=mc^2</m:t></m:r></m:oMath></a14:m>"


def _convert_paragraph(xml: str, *, title: bool = False) -> DoclingDocument:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0 if title else 6])
    shape = (
        slide.shapes.title
        if title
        else slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    )
    paragraph = shape.text_frame.paragraphs[0]._element
    paragraph.getparent().replace(
        paragraph, etree.fromstring(f"<a:p {_NAMESPACES}>{xml}</a:p>")
    )
    stream = BytesIO()
    presentation.save(stream)
    stream.seek(0)
    return (
        DocumentConverter(allowed_formats=[InputFormat.PPTX])
        .convert(DocumentStream(name="equations.pptx", stream=stream))
        .document
    )


@pytest.mark.parametrize("wrapped", [False, True])
def test_equation_only_shape_is_converted_once(wrapped: bool) -> None:
    xml = _EQUATION
    if wrapped:
        xml = (
            '<mc:AlternateContent><mc:Choice Requires="a14">'
            f"{xml}</mc:Choice><mc:Fallback>"
            "<a:r><a:t>Fallback equation</a:t></a:r>"
            "</mc:Fallback></mc:AlternateContent>"
        )

    doc = _convert_paragraph(xml)

    assert [(item.label, item.text) for item in doc.texts] == [
        (DocItemLabel.FORMULA, "E=mc^2")
    ]
    assert "E=mc^2" in doc.export_to_markdown()
    assert "Fallback" not in doc.export_to_markdown()
    assert doc.texts[0].prov[0].page_no == 1


@pytest.mark.parametrize("listed", [False, True])
def test_inline_equation_preserves_text_order_and_list_parent(listed: bool) -> None:
    properties = '<a:pPr><a:buChar char="•"/></a:pPr>' if listed else ""
    doc = _convert_paragraph(
        properties
        + "<a:r><a:t>Energy is </a:t></a:r>"
        + _EQUATION
        + "<a:r><a:t> in this model.</a:t></a:r>"
    )

    formulas = [item for item in doc.texts if item.label == DocItemLabel.FORMULA]
    assert [item.text for item in formulas] == ["E=mc^2"]
    markdown = doc.export_to_markdown()
    assert markdown.index("Energy") < markdown.index("E=mc^2") < markdown.index("model")
    if listed:
        assert any(isinstance(item, ListItem) for item in doc.texts)
        assert markdown.startswith("- ")


def test_unsupported_compatibility_choice_uses_fallback_text() -> None:
    doc = _convert_paragraph(
        '<mc:AlternateContent xmlns:future="https://example.com/unsupported">'
        '<mc:Choice Requires="future">'
        + _EQUATION
        + "</mc:Choice><mc:Fallback><a:r><a:t>Fallback text</a:t></a:r>"
        "</mc:Fallback></mc:AlternateContent>"
    )
    assert [item.text for item in doc.texts] == ["Fallback text"]


def test_fraction_uses_the_shared_omml_converter() -> None:
    doc = _convert_paragraph(
        "<a14:m><m:oMathPara><m:oMath><m:f>"
        "<m:num><m:r><m:t>a</m:t></m:r></m:num>"
        "<m:den><m:r><m:t>b</m:t></m:r></m:den>"
        "</m:f></m:oMath></m:oMathPara></a14:m>"
    )
    assert [(item.label, item.text) for item in doc.texts] == [
        (DocItemLabel.FORMULA, r"\frac{a}{b}")
    ]


def test_inline_equation_retains_the_title_label() -> None:
    doc = _convert_paragraph("<a:r><a:t>Energy </a:t></a:r>" + _EQUATION, title=True)
    title = next(item for item in doc.texts if item.label == DocItemLabel.TITLE)
    assert title.children
    assert "Energy" in doc.export_to_markdown()
    assert "E=mc^2" in doc.export_to_markdown()
