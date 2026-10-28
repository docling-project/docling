# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

import tracemalloc
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

from docling.backend.msexcel_backend import MsExcelDocumentBackend
from docling.datamodel.base_models import InputFormat
from docling.datamodel.document import InputDocument


def _load_backend_sheet(file_path: Path):
    in_doc = InputDocument(
        path_or_stream=file_path,
        format=InputFormat.XLSX,
        filename=file_path.stem,
        backend=MsExcelDocumentBackend,
    )
    backend = MsExcelDocumentBackend(in_doc=in_doc, path_or_stream=file_path)
    return backend, backend.workbook.active


def test_find_data_tables_does_not_materialize_empty_cells(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "first"
    sheet["AD33334"] = "last"  # ~1M-position bounding rectangle
    file_path = tmp_path / "sparse.xlsx"
    workbook.save(file_path)
    backend, loaded_sheet = _load_backend_sheet(file_path)
    cells_before = len(loaded_sheet._cells)

    tracemalloc.start()
    try:
        tables, _ = backend._find_data_tables(loaded_sheet)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert len(tables) == 2
    found = sorted((t.anchor, [(c.row, c.col, c.text) for c in t.data]) for t in tables)
    assert found == [
        ((0, 0), [(0, 0, "first")]),
        ((29, 33333), [(0, 0, "last")]),
    ]
    assert len(loaded_sheet._cells) == cells_before
    assert peak < 50 * 1024 * 1024


def test_sparse_merged_anchor_and_comment_only_cell_tables(tmp_path: Path) -> None:
    # Coverage test: passes on main too, pins output for merges and comments.
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "head"
    sheet["B1"] = "x"
    sheet["A2"] = "merged"
    sheet.merge_cells("A2:B2")
    sheet["C2"].comment = Comment("only a note", "Codex")
    sheet["A5000"] = "far"
    sheet["B5000"] = "far2"
    file_path = tmp_path / "sparse-merge-comment.xlsx"
    workbook.save(file_path)
    backend, loaded_sheet = _load_backend_sheet(file_path)

    tables, comment_map = backend._find_data_tables(loaded_sheet)

    found = sorted(
        (
            t.anchor,
            t.num_rows,
            t.num_cols,
            sorted((c.row, c.col, c.text, c.row_span, c.col_span) for c in t.data),
        )
        for t in tables
    )
    assert found == [
        (
            (0, 0),
            2,
            2,
            [(0, 0, "head", 1, 1), (0, 1, "x", 1, 1), (1, 0, "merged", 1, 2)],
        ),
        ((0, 4999), 1, 2, [(0, 0, "far", 1, 1), (0, 1, "far2", 1, 1)]),
    ]
    assert comment_map == {}  # comment-only cell outside any table is dropped


def test_missing_merge_anchor_inside_table_keeps_span(tmp_path: Path) -> None:
    # Coverage test for the contains_at fallback; main recreates the cell instead.
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "head"
    sheet.merge_cells("A2:B2")
    file_path = tmp_path / "missing-merge-anchor.xlsx"
    workbook.save(file_path)
    backend, loaded_sheet = _load_backend_sheet(file_path)
    del loaded_sheet._cells[(2, 1)]  # force the merged-anchor fallback

    tables, _ = backend._find_data_tables(loaded_sheet)

    assert len(tables) == 1
    table = tables[0]
    assert (table.num_rows, table.num_cols) == (2, 2)
    anchor = next(c for c in table.data if (c.row, c.col) == (1, 0))
    assert anchor.col_span == 2
