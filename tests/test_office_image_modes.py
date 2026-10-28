# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from zlib import crc32

import openpyxl
import pytest
from docx import Document
from docx.shared import Inches
from openpyxl.drawing.image import Image as XlImage
from PIL import Image, ImageCms, ImageFile, __version__ as PILLOW_VERSION
from pptx import Presentation

from docling.backend.msexcel_backend import MsExcelDocumentBackend
from docling.backend.msword_backend import MsWordDocumentBackend
from docling.backend.utils.image import normalize_image_for_png
from docling.datamodel.base_models import ConversionStatus
from docling.document_converter import DocumentConverter


def _cmyk_jpeg(profile: bytes | None = None) -> BytesIO:
    jpeg = BytesIO()
    Image.new("CMYK", (200, 100), (0, 255, 255, 0)).save(
        jpeg, format="JPEG", icc_profile=profile
    )
    jpeg.seek(0)
    return jpeg


def _save_xlsx_with_picture(path: Path, picture: BytesIO) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "label"
    sheet.add_image(XlImage(picture), "B2")
    book.save(path)


@pytest.mark.parametrize("profile", [None, b"invalid ICC profile"])
@pytest.mark.parametrize("suffix", ["pptx", "docx", "xlsx"])
def test_cmyk_office_picture_without_libreoffice(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, suffix: str, profile: bytes | None
) -> None:
    # Exercise the real unavailable-converter path without changing the system.
    monkeypatch.setenv("PATH", str(tmp_path))
    monkeypatch.delenv("DOCLING_LIBREOFFICE_CMD", raising=False)
    monkeypatch.setattr(
        "docling.backend.msword_backend.get_docx_to_pdf_converter", lambda: None
    )
    monkeypatch.setattr(
        "docling.backend.msexcel_backend.get_docx_to_pdf_converter", lambda: None
    )
    jpeg = _cmyk_jpeg(profile)
    path = tmp_path / f"cmyk.{suffix}"
    if suffix == "pptx":
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(jpeg, Inches(1), Inches(1))
        deck.save(path)
    elif suffix == "xlsx":
        _save_xlsx_with_picture(path, jpeg)
    else:
        # python-docx rejects CMYK JPEG headers, so replace the embedded RGB JPEG.
        rgb = BytesIO()
        Image.new("RGB", (200, 100), "red").save(rgb, format="JPEG")
        rgb.seek(0)
        source = Document()
        source.add_picture(rgb, width=Inches(2))
        template = BytesIO()
        source.save(template)
        template.seek(0)
        with ZipFile(template) as src, ZipFile(path, "w") as dst:
            for entry in src.infolist():
                data = (
                    jpeg.getvalue()
                    if entry.filename.startswith("word/media/")
                    else src.read(entry.filename)
                )
                dst.writestr(entry, data)

    result = DocumentConverter().convert(path)
    assert result.status == ConversionStatus.SUCCESS
    assert len(result.document.pictures) == 1
    image_ref = result.document.pictures[0].image
    assert image_ref is not None
    image = image_ref.pil_image
    assert image is not None
    assert image.size == (200, 100)
    assert image.convert("RGB").getpixel((0, 0)) == (255, 0, 0)


def test_cmyk_xlsx_picture_is_not_rendered_by_libreoffice(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Stand in for an installed LibreOffice: it would re-render the picture at
    # another size, so the stored picture must not be routed to it.
    monkeypatch.setattr(
        MsExcelDocumentBackend,
        "_convert_emf_to_pil",
        lambda self, image_bytes: Image.new("RGB", (401, 201), "white"),
    )
    path = tmp_path / "cmyk.xlsx"
    _save_xlsx_with_picture(path, _cmyk_jpeg())

    result = DocumentConverter().convert(path)
    assert result.status == ConversionStatus.SUCCESS
    assert len(result.document.pictures) == 1
    image_ref = result.document.pictures[0].image
    assert image_ref is not None
    image = image_ref.pil_image
    assert image is not None
    assert image.size == (200, 100)


@pytest.mark.parametrize(
    "chunk_type, payload",
    [
        pytest.param(
            b"cHRM",
            b"",
            marks=pytest.mark.skipif(
                int(PILLOW_VERSION.split(".")[0]) < 13,
                reason="Pillow before 13 accepts truncated cHRM chunks",
            ),
        ),
        (b"gAMA", b"\x00\x00"),
        (b"tRNS", b"\x00\x00"),
    ],
)
@pytest.mark.parametrize("render_available", [True, False])
def test_docx_truncated_png_uses_picture_fallback(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    chunk_type: bytes,
    payload: bytes,
    render_available: bool,
) -> None:
    monkeypatch.setattr(ImageFile, "LOAD_TRUNCATED_IMAGES", False)
    rendered_elements: list[object] = []

    def render(
        self: MsWordDocumentBackend,
        elements: object,
        element_tag: str | list[str] | None = None,
    ) -> Image.Image | None:
        assert element_tag == ["drawing", "pict"]
        rendered_elements.append(elements)
        return Image.new("RGB", (240, 120), "green") if render_available else None

    monkeypatch.setattr(MsWordDocumentBackend, "_convert_elements_via_docx", render)
    png = BytesIO()
    Image.new("RGB", (200, 100), "red").save(png, format="PNG")
    png.seek(0)
    source = Document()
    source.add_paragraph("Before damaged picture")
    source.add_picture(png, width=Inches(2))
    source.add_paragraph("After damaged picture")
    healthy_png = BytesIO()
    Image.new("RGB", (200, 100), "blue").save(healthy_png, format="PNG")
    healthy_png.seek(0)
    source.add_picture(healthy_png, width=Inches(2))
    template = BytesIO()
    source.save(template)
    template.seek(0)

    # Insert a short ancillary chunk with a valid CRC after the PNG IHDR chunk.
    chunk = len(payload).to_bytes(4, "big") + chunk_type + payload
    chunk += crc32(chunk_type + payload).to_bytes(4, "big")
    png_bytes = png.getvalue()
    damaged_png = png_bytes[:33] + chunk + png_bytes[33:]
    path = tmp_path / "truncated.docx"
    with ZipFile(template) as src, ZipFile(path, "w") as dst:
        for entry in src.infolist():
            data = (
                damaged_png
                if entry.filename == "word/media/image1.png"
                else src.read(entry.filename)
            )
            dst.writestr(entry, data)

    result = DocumentConverter().convert(path)
    assert result.status == ConversionStatus.SUCCESS
    assert len(rendered_elements) == 1
    doc = result.document
    assert [text.text for text in doc.texts] == [
        "Before damaged picture",
        "After damaged picture",
    ]
    assert [ref.cref for ref in doc.body.children] == [
        "#/texts/0",
        "#/pictures/0",
        "#/texts/1",
        "#/pictures/1",
    ]
    assert len(doc.pictures) == 2
    image_ref = doc.pictures[0].image
    if render_available:
        assert image_ref is not None
        image = image_ref.pil_image
        assert image is not None
        assert image.size == (240, 120)
        assert image.getpixel((0, 0)) == (0, 128, 0)
    else:
        assert image_ref is None
    healthy_ref = doc.pictures[1].image
    assert healthy_ref is not None
    healthy_image = healthy_ref.pil_image
    assert healthy_image is not None
    assert healthy_image.getpixel((0, 0)) == (0, 0, 255)


def test_embedded_icc_profile_is_used() -> None:
    # LAB also requires conversion for PNG and has a profile Pillow can generate.
    image = Image.new("LAB", (2, 2), (128, 180, 80))
    profile = ImageCms.ImageCmsProfile(ImageCms.createProfile("LAB"))
    image.info["icc_profile"] = profile.tobytes()
    expected = ImageCms.profileToProfile(
        image, profile, ImageCms.createProfile("sRGB"), outputMode="RGB"
    )
    actual = normalize_image_for_png(image)
    assert actual.mode == "RGB"
    assert actual.tobytes() == expected.tobytes()


def test_pa_preserves_per_pixel_alpha_in_png() -> None:
    image = Image.new("PA", (2, 1))
    image.putpalette([10, 20, 30] + [0, 0, 0] * 255)
    image.putpixel((0, 0), (0, 40))
    image.putpixel((1, 0), (0, 200))

    buffer = BytesIO()
    normalize_image_for_png(image).save(buffer, format="PNG")
    buffer.seek(0)
    with Image.open(buffer) as decoded:
        assert decoded.mode == "RGBA"
        assert decoded.getpixel((0, 0)) == (10, 20, 30, 40)
        assert decoded.getpixel((1, 0)) == (10, 20, 30, 200)


def test_integer_image_preserves_png_values() -> None:
    image = Image.new("I", (6, 1))
    image.putdata([-1, 0, 255, 256, 65535, 65536])
    buffer = BytesIO()

    normalize_image_for_png(image).save(buffer, format="PNG")

    buffer.seek(0)
    with Image.open(buffer) as decoded:
        assert [decoded.getpixel((x, 0)) for x in range(6)] == [
            0,
            0,
            255,
            256,
            65535,
            65535,
        ]


@pytest.mark.parametrize(
    "mode, color", [("RGB", (10, 20, 30)), ("RGBA", (10, 20, 30, 40))]
)
def test_compatible_pixels_survive_png_serialization(mode: str, color: tuple) -> None:
    image = Image.new(mode, (2, 2), color)
    buffer = BytesIO()
    normalize_image_for_png(image).save(buffer, format="PNG")
    buffer.seek(0)
    with Image.open(buffer) as decoded:
        assert decoded.mode == mode
        assert decoded.getpixel((0, 0)) == color
