# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for Word style resolution.

Kept separate from ``test_backend_msword.py`` so that file stays under the
repository's per-file line limit.
"""

from docling_core.types.doc import DoclingDocument
from docx import Document

from docling.backend.msword_backend import MsWordDocumentBackend
from docling.datamodel.base_models import InputFormat
from docling.datamodel.document import InputDocument


def _convert_docx(document, docx_path) -> DoclingDocument:
    document.save(str(docx_path))
    in_doc = InputDocument(
        path_or_stream=docx_path,
        format=InputFormat.DOCX,
        backend=MsWordDocumentBackend,
        filename=docx_path.name,
    )
    return MsWordDocumentBackend(in_doc=in_doc, path_or_stream=docx_path).convert()


def test_style_lookups_do_not_leak_between_documents(tmp_path):
    """Memoized style lookups must be scoped to the document they came from.

    Both paragraphs resolve to the default paragraph style through the same
    ``get_style(None, PARAGRAPH)`` lookup, but only the first document makes
    that style bold. Converting them one after the other in one process must
    report the bold of each document's own style.
    """
    bold_doc = Document()
    bold_doc.styles["Normal"].font.bold = True
    bold_doc.add_paragraph("bold by style")
    plain_doc = Document()
    plain_doc.add_paragraph("plain")

    bold_texts = _convert_docx(bold_doc, tmp_path / "bold.docx").texts
    plain_texts = _convert_docx(plain_doc, tmp_path / "plain.docx").texts

    assert [t.text for t in bold_texts] == ["bold by style"]
    assert bold_texts[0].formatting is not None
    assert bold_texts[0].formatting.bold is True
    assert [t.text for t in plain_texts] == ["plain"]
    assert plain_texts[0].formatting is not None
    assert plain_texts[0].formatting.bold is False
