# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for video and audio shapes on PowerPoint slides."""

import zipfile
from pathlib import Path

from docling_core.types.doc import DoclingDocument, PictureItem

from .test_backend_pptx import get_converter

_P14_MEDIA = "{http://schemas.microsoft.com/office/powerpoint/2010/main}media"


def _add_media(tmp_path: Path, file_name: str, mime_type: str):
    """Add a media shape the way PowerPoint stores it, and return the deck and shape.

    python-pptx writes the same XML as PowerPoint: an ``a:videoFile`` link plus a
    ``p14:media`` extension, both pointing at one part under ``ppt/media``.
    """
    from pptx import Presentation
    from pptx.util import Inches

    media_path = tmp_path / file_name
    media_path.write_bytes(b"not decoded by the backend")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
    title.text_frame.text = "Slide with media"
    shape = slide.shapes.add_movie(
        str(media_path),
        Inches(1),
        Inches(2),
        Inches(4),
        Inches(2.25),
        mime_type=mime_type,
    )
    return prs, slide, shape


def _media_meta(picture: PictureItem) -> dict:
    assert picture.meta is not None
    return picture.meta.get_custom_part()


def _convert(prs, path: Path) -> DoclingDocument:
    prs.save(path)
    return get_converter().convert(path).document


def test_pptx_embedded_video_is_a_picture_with_its_media_path(tmp_path: Path):
    """An embedded video becomes a picture of its poster frame on its slide.

    The picture records the path of the video inside the package, so the video
    can be read from the deck, e.g. to transcribe it.
    """
    prs, _, _ = _add_media(tmp_path, "clip.mp4", "video/mp4")
    deck = tmp_path / "video.pptx"
    doc = _convert(prs, deck)

    assert len(doc.pictures) == 1
    picture = doc.pictures[0]
    assert picture.image is not None
    assert picture.prov[0].page_no == 1
    assert _media_meta(picture) == {"docling__video": "ppt/media/media1.mp4"}
    with zipfile.ZipFile(deck) as package:
        assert package.read("ppt/media/media1.mp4") == b"not decoded by the backend"
    assert doc.export_to_markdown() == (
        "Slide with media\n\n<!-- image -->\n\nppt/media/media1.mp4"
    )


def test_pptx_audio_picture_records_its_media_path(tmp_path: Path):
    """An audio shape keeps its speaker icon and now says which file it plays.

    PowerPoint stores audio like video, with ``a:audioFile`` in place of
    ``a:videoFile``; python-pptx reads it as a plain picture.
    """
    from pptx.oxml.ns import qn

    prs, _, shape = _add_media(tmp_path, "tone.mp3", "audio/mpeg")
    shape._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile")[0].tag = qn("a:audioFile")
    doc = _convert(prs, tmp_path / "audio.pptx")

    assert len(doc.pictures) == 1
    assert doc.pictures[0].image is not None
    assert _media_meta(doc.pictures[0]) == {"docling__audio": "ppt/media/media1.mp3"}


def test_pptx_linked_video_records_its_link(tmp_path: Path):
    """A video that is linked, not embedded, records the link target."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn

    prs, slide, shape = _add_media(tmp_path, "clip.mp4", "video/mp4")
    link = "https://example.com/talk.mp4"
    video_rid = slide.part.relate_to(link, RT.VIDEO, is_external=True)
    media_rid = slide.part.relate_to(link, RT.MEDIA, is_external=True)
    shape._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile")[0].set(
        qn("r:link"), video_rid
    )
    p14_media = next(shape._element.iter(_P14_MEDIA))
    del p14_media.attrib[qn("r:embed")]
    p14_media.set(qn("r:link"), media_rid)
    doc = _convert(prs, tmp_path / "linked.pptx")

    assert len(doc.pictures) == 1
    assert _media_meta(doc.pictures[0]) == {"docling__video": link}


def test_pptx_picture_without_media_has_no_media_meta(tmp_path: Path):
    """A plain picture is unchanged: it carries no media field."""
    from io import BytesIO

    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    image = BytesIO()
    Image.new("RGB", (40, 20), "red").save(image, format="PNG")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(image, Inches(1), Inches(1))
    doc = _convert(prs, tmp_path / "picture.pptx")

    assert len(doc.pictures) == 1
    assert doc.pictures[0].meta is None
