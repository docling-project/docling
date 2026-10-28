# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for a hyperlink attached to a whole PowerPoint shape."""

from pathlib import Path

from docling_core.types.doc import ContentLayer, TextItem
from pptx import Presentation
from pptx.util import Inches

from docling.datamodel.document import DoclingDocument

from .test_backend_pptx import get_converter


def _links_by_text(doc: DoclingDocument) -> dict[str, str | None]:
    return {
        item.text: None if item.hyperlink is None else str(item.hyperlink)
        for item, _ in doc.iterate_items(included_content_layers=set(ContentLayer))
        if isinstance(item, TextItem)
    }


def _convert(prs, path: Path) -> DoclingDocument:
    prs.save(path)
    return get_converter().convert(path).document


def test_pptx_shape_hyperlink_reaches_its_text(tmp_path: Path):
    """A link on a whole shape (``a:hlinkClick`` under ``p:cNvPr``) is kept.

    Every paragraph of the shape carries the link; a shape without one is
    left alone.
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    plain = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    plain.text_frame.text = "No link here"
    button = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    button.text_frame.text = "Watch the talk"
    button.text_frame.add_paragraph().text = "Second line"
    button.click_action.hyperlink.address = "https://example.com/talk"

    doc = _convert(prs, tmp_path / "shape_link.pptx")

    assert _links_by_text(doc) == {
        "No link here": None,
        "Watch the talk": "https://example.com/talk",
        "Second line": "https://example.com/talk",
    }
    assert "[Watch the talk](https://example.com/talk)" in doc.export_to_markdown()


def test_pptx_shape_hyperlink_reaches_list_items(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    body = slide.placeholders[1]
    body.text_frame.text = "First point"
    body.text_frame.add_paragraph().text = "Second point"
    body.click_action.hyperlink.address = "https://example.com/points"

    doc = _convert(prs, tmp_path / "list_link.pptx")

    links = _links_by_text(doc)
    assert links["First point"] == "https://example.com/points"
    assert links["Second point"] == "https://example.com/points"


def test_pptx_shape_relative_link_is_a_path(tmp_path: Path):
    """An address without a URL scheme is kept as a path, as the DOCX backend does."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Open the appendix"
    box.click_action.hyperlink.address = "appendix/notes.docx"

    doc = _convert(prs, tmp_path / "relative_link.pptx")

    (item,) = [t for t in doc.texts if t.text == "Open the appendix"]
    assert item.hyperlink == Path("appendix/notes.docx")


def test_pptx_slide_jump_action_is_not_a_hyperlink(tmp_path: Path):
    """A click action that jumps to another slide has no address to keep."""
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[6])
    second = prs.slides.add_slide(prs.slide_layouts[6])
    box = first.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Go to the next slide"
    box.click_action.target_slide = second

    doc = _convert(prs, tmp_path / "slide_jump.pptx")

    assert _links_by_text(doc) == {"Go to the next slide": None}
