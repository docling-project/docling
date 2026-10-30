# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Slide video and audio converted with the real audio and video pipelines."""

import shutil
from pathlib import Path

import pytest
from docling_core.types.doc import GroupItem, PictureItem, TextItem

from docling.datamodel.base_models import ConversionStatus, InputFormat
from docling.datamodel.pipeline_options import ConvertPipelineOptions
from docling.document_converter import DocumentConverter, PowerpointFormatOption

pytestmark = pytest.mark.ml_asr

_SOURCES = Path("./tests/data/audio/sources")


@pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="ffmpeg not available")
def test_slide_video_and_audio_are_transcribed(tmp_path: Path):
    """The transcript of each media file follows its picture on the slide.

    A video also gives its sampled frames.
    """
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_movie(
        str(_SOURCES / "sample_10s_video-mp4.mp4"),
        Inches(1),
        Inches(1),
        Inches(4),
        Inches(2.25),
        mime_type="video/mp4",
    )
    audio = slide.shapes.add_movie(
        str(_SOURCES / "sample_10s.mp3"),
        Inches(6),
        Inches(1),
        Inches(1),
        Inches(1),
        mime_type="audio/mpeg",
    )
    audio._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile")[0].tag = qn("a:audioFile")
    deck = tmp_path / "media.pptx"
    prs.save(deck)

    converter = DocumentConverter(
        allowed_formats=[InputFormat.PPTX],
        format_options={
            InputFormat.PPTX: PowerpointFormatOption(
                pipeline_options=ConvertPipelineOptions(do_media_conversion=True)
            )
        },
    )
    result = converter.convert(deck)

    assert result.status == ConversionStatus.SUCCESS
    doc = result.document
    media = [item for item in doc.groups if item.name.startswith("media: ")]
    assert [group.name for group in media] == [
        "media: ppt/media/media1.mp4",
        "media: ppt/media/media2.mp3",
    ]
    for group in media:
        items = [ref.resolve(doc) for ref in group.children]
        texts = [item for item in items if isinstance(item, TextItem)]
        assert any(text.text.strip() for text in texts)
        assert all(text.source and text.prov[0].page_no == 1 for text in texts)
    video_items = [ref.resolve(doc) for ref in media[0].children]
    assert any(isinstance(item, PictureItem) for item in video_items)
    assert not any(isinstance(item, GroupItem) for item in video_items)


@pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="ffmpeg not available")
def test_silent_slide_video_keeps_its_frames(tmp_path: Path):
    """A screen recording without audio is no error: its frames come without cues of speech."""
    import subprocess

    from pptx import Presentation
    from pptx.util import Inches

    video = tmp_path / "silent.mp4"
    subprocess.run(
        [
            "ffmpeg",
            "-hide_banner",
            "-loglevel",
            "error",
            "-f",
            "lavfi",
            "-i",
            "testsrc=size=160x90:rate=10:duration=3",
            "-c:v",
            "mpeg4",
            str(video),
        ],
        check=True,
    )
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_movie(
        str(video), Inches(1), Inches(1), Inches(4), Inches(2.25), mime_type="video/mp4"
    )
    deck = tmp_path / "silent.pptx"
    prs.save(deck)

    converter = DocumentConverter(
        allowed_formats=[InputFormat.PPTX],
        format_options={
            InputFormat.PPTX: PowerpointFormatOption(
                pipeline_options=ConvertPipelineOptions(do_media_conversion=True)
            )
        },
    )
    result = converter.convert(deck)

    assert result.status == ConversionStatus.SUCCESS
    assert result.errors == []
    media = [item for item in result.document.groups if item.name.startswith("media: ")]
    assert len(media) == 1
    items = [ref.resolve(result.document) for ref in media[0].children]
    assert items and all(isinstance(item, PictureItem) for item in items)
