# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

from io import BytesIO

import pytest
from docling_core.types.doc import ContentLayer
from openpyxl import Workbook

from docling.datamodel.base_models import DocumentStream, InputFormat
from docling.document_converter import DocumentConverter


def _convert(workbook: Workbook):
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return (
        DocumentConverter(allowed_formats=[InputFormat.XLSX])
        .convert(DocumentStream(name="hidden.xlsx", stream=stream))
        .document
    )


def test_hidden_rows_and_columns_stay_out_of_visible_table() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Price", "Internal cost"])
    sheet.append(["Orange juice", 12.5, 9.1])
    sheet.append(["Discontinued", 7, 5])
    sheet.append(["Apple juice", 7.25, 5.2])
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions["C"].hidden = True

    doc = _convert(workbook)

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert (table.data.num_rows, table.data.num_cols) == (3, 2)
    assert [cell.text for cell in table.data.table_cells] == [
        "Item",
        "Price",
        "Orange juice",
        "12.5",
        "Apple juice",
        "7.25",
    ]
    assert "Discontinued" not in doc.export_to_markdown()
    assert "Internal cost" not in doc.export_to_markdown()
    hidden = [
        item for item in doc.texts if item.content_layer == ContentLayer.INVISIBLE
    ]
    assert {item.text for item in hidden} == {
        "Internal cost",
        "9.1",
        "Discontinued",
        "7",
        "5",
        "5.2",
    }
    assert table.prov[0].bbox.as_tuple() == (0, 0, 3, 4)


def test_grouped_hidden_columns_and_merged_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Private 1", "Private 2", "Value"])
    sheet.append(["Spanning value", None, None, None])
    sheet.merge_cells("A2:D2")
    sheet.column_dimensions.group("B", "C", hidden=True)

    doc = _convert(workbook)

    table = doc.tables[0]
    assert (table.data.num_rows, table.data.num_cols) == (2, 2)
    assert [cell.text for cell in table.data.table_cells] == [
        "Name",
        "Value",
        "Spanning value",
    ]
    assert table.data.table_cells[-1].col_span == 2
    assert "Private" not in doc.export_to_markdown()


@pytest.mark.parametrize("axis", ["row", "column"])
def test_fully_hidden_table_keeps_content_in_invisible_layer(axis: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Hidden value"])
    if axis == "row":
        sheet.row_dimensions[1].hidden = True
    else:
        sheet.column_dimensions["A"].hidden = True

    doc = _convert(workbook)

    assert not doc.tables
    assert "Hidden value" not in doc.export_to_markdown()
    assert [(item.text, item.content_layer) for item in doc.texts] == [
        ("Hidden value", ContentLayer.INVISIBLE)
    ]


def test_hidden_sheet_keeps_its_original_table() -> None:
    workbook = Workbook()
    workbook.active.append(["Visible"])
    sheet = workbook.create_sheet("Hidden")
    sheet.append(["Secret", "Value"])
    sheet.append(["Private", 42])
    sheet.row_dimensions[2].hidden = True
    sheet.sheet_state = "hidden"

    doc = _convert(workbook)

    table = doc.tables[1]
    assert table.content_layer == ContentLayer.INVISIBLE
    assert (table.data.num_rows, table.data.num_cols) == (2, 2)
    assert [cell.text for cell in table.data.table_cells] == [
        "Secret",
        "Value",
        "Private",
        "42",
    ]


def test_hidden_section_label_and_merged_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Private title", None])
    sheet.merge_cells("A1:B1")
    sheet.append(["Name", "Value"])
    sheet.append(["Shared label", "Hidden value"])
    sheet.append([None, "Visible value"])
    sheet.merge_cells("A3:A4")
    sheet.row_dimensions[1].hidden = True
    sheet.row_dimensions[3].hidden = True

    doc = _convert(workbook)

    assert "Private title" not in doc.export_to_markdown()
    assert "Hidden value" not in doc.export_to_markdown()
    table = doc.tables[0]
    assert (table.data.num_rows, table.data.num_cols) == (2, 2)
    assert [cell.text for cell in table.data.table_cells] == [
        "Name",
        "Value",
        "Shared label",
        "Visible value",
    ]
    assert table.data.table_cells[2].row_span == 1
    assert table.prov[0].bbox.as_tuple() == (0, 1, 2, 4)


def test_hidden_dimensions_with_a_table_offset() -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row, values in enumerate(
        [
            ["Name", "Internal 1", "Internal 2", "Value"],
            ["Old item", "secret", "secret", 10],
            ["Current item", "secret", "secret", 20],
        ],
        start=5,
    ):
        for column, value in enumerate(values, start=6):
            sheet.cell(row=row, column=column, value=value)
    sheet.column_dimensions.group("A", "B", hidden=True)
    sheet.column_dimensions.group("G", "H", hidden=True)
    sheet.row_dimensions[6].hidden = True

    doc = _convert(workbook)

    table = doc.tables[0]
    assert (table.data.num_rows, table.data.num_cols) == (2, 2)
    assert [cell.text for cell in table.data.table_cells] == [
        "Name",
        "Value",
        "Current item",
        "20",
    ]
    assert table.prov[0].bbox.as_tuple() == (5, 4, 9, 7)
    assert "secret" not in doc.export_to_markdown()
    assert "Old item" not in doc.export_to_markdown()
