# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for Excel combination charts: several plots in one chart."""

import datetime
import logging
from pathlib import Path
from typing import Optional

import pytest
from docling_core.types.doc import PictureClassificationLabel, PictureItem
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import StrData, StrRef, StrVal
from openpyxl.chart.text import Text
from openpyxl.chart.title import Title

from docling.backend import msexcel_backend
from docling.datamodel.backend_options import MsExcelBackendOptions
from docling.datamodel.base_models import InputFormat
from docling.document_converter import DocumentConverter, ExcelFormatOption

from .test_backend_msexcel import get_converter

_MONTHS = [datetime.datetime(2012, month, 1) for month in range(1, 13)]
_BARS = {
    "East": (85, 34, 100, 80, 116, 88, 71, 90, 121, 110, 101, 81),
    "West": (94, 72, 62, 91, 99, 70, 71, 46, 91, 57, 32, 157),
    "North": (91, 59, 85, 67, 130, 66, 51, 40, 62, 61, 61, 52),
    "South": (71, 110, 101, 28, 39, 52, 48, 125, 43, 122, 60, 63),
}
_LINES = {
    "Margin": tuple(
        n / 1000 for n in (126, 74, 110, 92, 86, 89, 77, 40, 101, 62, 81, 62)
    ),
    "Target": tuple(
        n / 1000 for n in (91, 75, 94, 72, 103, 75, 66, 81, 85, 94, 69, 95)
    ),
}


def _save_combo_chart_workbook(
    path: Path,
    time_of_day: datetime.time = datetime.time(0),
    linked_axis_title: bool = False,
    range_end_row: Optional[int] = None,
) -> None:
    """Save a workbook whose chart draws columns and lines on a second axis.

    This is how Excel stores a "Clustered Column - Line on Secondary Axis"
    chart: a ``c:barChart`` with the four region series, then a
    ``c:lineChart`` with the two ratio series on a secondary value axis. Both
    plots share the date categories in column A, which have ``time_of_day``.
    The category axis is titled "Region Position" and the primary value axis
    "Bar Sales". With ``linked_axis_title`` the category axis title is linked
    to cell A1, so its text is in a cell cache instead of in rich text. The
    ranges of the chart end at ``range_end_row``, or at the last data row.
    """
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.append(["Month", *_BARS, *_LINES])
    for row, month in enumerate(_MONTHS):
        sheet.append(
            [datetime.datetime.combine(month, time_of_day)]
            + [v[row] for v in _BARS.values()]
            + [v[row] for v in _LINES.values()]
        )
    last_row = range_end_row or 1 + len(_MONTHS)
    categories = Reference(sheet, min_col=1, min_row=2, max_row=last_row)

    bars = BarChart()
    bars.title = "Sales and Margin"
    bars.add_data(
        Reference(sheet, min_col=2, max_col=5, min_row=1, max_row=last_row),
        titles_from_data=True,
    )
    bars.set_categories(categories)
    if linked_axis_title:
        cache = StrData(ptCount=1, pt=[StrVal(idx=0, v="Region Position")])
        bars.x_axis.title = Title(
            tx=Text(strRef=StrRef(f="'Sheet'!$A$1", strCache=cache))
        )
    else:
        bars.x_axis.title = "Region Position"
    bars.y_axis.title = "Bar Sales"

    lines = LineChart()
    lines.add_data(
        Reference(sheet, min_col=6, max_col=7, min_row=1, max_row=last_row),
        titles_from_data=True,
    )
    lines.set_categories(categories)
    lines.y_axis.axId = 200
    lines.y_axis.crosses = "max"

    bars += lines
    sheet.add_chart(bars, "J2")
    workbook.save(path)


def _only_chart(path: Path, converter: Optional[DocumentConverter] = None):
    doc = (converter or get_converter()).convert(path).document
    pictures = list(doc.pictures)
    assert len(pictures) == 1
    return pictures[0]


def _grid(picture: PictureItem) -> list[list[str]]:
    assert picture.meta is not None and picture.meta.tabular_chart is not None
    return [
        [cell.text for cell in row]
        for row in picture.meta.tabular_chart.chart_data.grid
    ]


def _labels(picture: PictureItem) -> list[str]:
    assert picture.meta is not None and picture.meta.classification is not None
    return [p.class_name for p in picture.meta.classification.predictions]


def test_combo_chart_lists_each_plot_type_and_keeps_all_series(tmp_path: Path):
    """A column + line chart keeps its line series and is classified as both.

    openpyxl gives the first plot as the chart, so before only the four column
    series were read and the chart was a bar chart only. The main prediction
    stays the column plot, so serializers that read only the main prediction
    still treat the picture as a bar chart.
    """
    path = tmp_path / "combo_chart.xlsx"
    _save_combo_chart_workbook(path)

    picture = _only_chart(path)

    assert _labels(picture) == [
        PictureClassificationLabel.BAR_CHART.value,
        PictureClassificationLabel.LINE_CHART.value,
    ]
    grid = _grid(picture)
    assert grid[0][1:] == ["East", "West", "North", "South", "Margin", "Target"]
    assert grid[1][1:] == ["85", "94", "91", "71", "0.126", "0.091"]
    assert len(grid) == 1 + len(_MONTHS)


@pytest.mark.parametrize(
    ("time_of_day", "suffix"),
    [(datetime.time(0), ""), (datetime.time(12, 30), "T12:30:00")],
    ids=["date", "date-and-time"],
)
def test_date_categories_are_iso_dates(
    tmp_path: Path, time_of_day: datetime.time, suffix: str
):
    """Date categories read as ISO 8601, with a time of day only when they have one.

    openpyxl reads a date cell as a ``datetime``, which was written as
    ``2012-01-01 00:00:00``.
    """
    path = tmp_path / "combo_chart.xlsx"
    _save_combo_chart_workbook(path, time_of_day=time_of_day)

    grid = _grid(_only_chart(path))

    assert [row[0] for row in grid[1:]] == [
        month.date().isoformat() + suffix for month in _MONTHS
    ]


@pytest.mark.parametrize("linked", [False, True], ids=["rich-text", "cell-linked"])
def test_category_axis_title_heads_the_category_column(tmp_path: Path, linked: bool):
    """The category axis title fills the top-left cell of the data table.

    That cell was always empty. A title linked to a cell is read from the cell
    cache. The value axis title has no column to head, so it stays out.
    """
    path = tmp_path / "combo_chart.xlsx"
    _save_combo_chart_workbook(path, linked_axis_title=linked)

    grid = _grid(_only_chart(path))

    assert grid[0] == ["Region Position", *_BARS, *_LINES]


def test_render_workbook_holds_the_data_of_every_plot(tmp_path: Path, monkeypatch):
    """The workbook that LibreOffice draws holds the line series data too.

    That workbook keeps every plot of the chart, but before only the data of
    the first plot was copied into it, so the line plot was drawn without data.
    A stand-in for LibreOffice keeps the workbook it gets, so this runs without
    LibreOffice.
    """
    path = tmp_path / "combo_chart.xlsx"
    _save_combo_chart_workbook(path)
    rendered: list[Workbook] = []

    def libreoffice(input_path: Path, output_path: Path) -> None:
        rendered.append(load_workbook(input_path))

    monkeypatch.setattr(
        msexcel_backend, "get_docx_to_pdf_converter", lambda: libreoffice
    )
    options = MsExcelBackendOptions(render_chart_images=True)
    converter = DocumentConverter(
        allowed_formats=[InputFormat.XLSX],
        format_options={InputFormat.XLSX: ExcelFormatOption(backend_options=options)},
    )

    _only_chart(path, converter)

    assert len(rendered) == 1
    data_sheet = rendered[0]["Sheet"]
    assert [data_sheet.cell(row=1, column=col).value for col in (6, 7)] == list(_LINES)
    assert data_sheet.cell(row=2, column=6).value == _LINES["Margin"][0]


def test_ranges_far_below_the_data_read_only_the_used_rows(tmp_path: Path, caplog):
    """Ranges that end far below the data stop at the last used row.

    Every range here ends at row 1,000,000. Before, the backend read every
    cell of every range, and a 6 KB file took minutes and gigabytes of memory
    to give a table of 1,000,000 rows.
    """
    path = tmp_path / "combo_chart.xlsx"
    _save_combo_chart_workbook(path, range_end_row=1_000_000)

    with caplog.at_level(logging.WARNING, logger="docling.backend.msexcel_backend"):
        grid = _grid(_only_chart(path))

    assert len(grid) == 1 + len(_MONTHS)
    assert "Reading only" not in caplog.text


def test_chart_reads_at_most_the_cell_cap(tmp_path: Path, monkeypatch, caplog):
    """The ranges of a chart share the cell cap, so the chart reads at most the cap.

    The categories and the 6 series each get a seventh of the cap. A cap of 70
    cells stands in for the real one, so each range reads 10 of its 12 rows.
    """
    monkeypatch.setattr(msexcel_backend, "MAX_CHART_TABLE_CELLS", 70)
    path = tmp_path / "combo_chart.xlsx"
    _save_combo_chart_workbook(path)

    with caplog.at_level(logging.WARNING, logger="docling.backend.msexcel_backend"):
        grid = _grid(_only_chart(path))

    assert len(grid) == 1 + 10
    assert grid[10][0] == _MONTHS[9].date().isoformat()
    assert "Reading only the first 10 of 12 rows" in caplog.text


def test_whole_column_range_ends_at_the_last_used_row(tmp_path: Path):
    """A series that refers to whole columns reads the rows that the sheet uses.

    A range such as ``$B:$B`` has no end row. Before, reading it raised an
    error, and the chart was lost.
    """
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    for month, value in (("Jan", 10), ("Feb", 20), ("Mar", 30)):
        sheet.append([month, value])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3))
    chart.set_categories(Reference(sheet, min_col=1, min_row=1, max_row=3))
    chart.series[0].val.numRef.f = "'Sheet'!$B:$B"
    chart.series[0].cat.numRef.f = "'Sheet'!$A:$A"
    sheet.add_chart(chart, "D2")
    path = tmp_path / "whole_columns.xlsx"
    workbook.save(path)

    grid = _grid(_only_chart(path))

    assert grid[1:] == [["Jan", "10"], ["Feb", "20"], ["Mar", "30"]]
