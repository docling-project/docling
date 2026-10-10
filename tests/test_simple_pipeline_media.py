# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for the conversion of slide video and audio in the simple pipeline.

A stub stands in for the audio and video pipelines, so these tests need no
speech model. ``test_simple_pipeline_media_asr.py`` runs the real models.
"""

from pathlib import Path
from typing import Optional

import pytest
from docling_core.types.doc import (
    ContentLayer,
    DocItemLabel,
    DoclingDocument,
    GroupItem,
    GroupLabel,
    PictureItem,
    TextItem,
    TrackSource,
)

from docling.datamodel.backend_options import MsPowerpointBackendOptions
from docling.datamodel.base_models import (
    ConversionStatus,
    DoclingComponentType,
    ErrorItem,
    InputFormat,
)
from docling.datamodel.document import ConversionResult, InputDocument
from docling.datamodel.pipeline_options import ConvertPipelineOptions
from docling.document_converter import DocumentConverter, PowerpointFormatOption
from docling.pipeline.simple_pipeline import SimplePipeline

_MEDIA_BYTES = b"bytes of the media file"
_TRANSCRIPT = "Spoken words"


class _StubMediaPipeline:
    """Stands in for the audio and video pipelines: no model, a fixed transcript."""

    def __init__(self, status: ConversionStatus = ConversionStatus.SUCCESS) -> None:
        self.status = status
        self.calls: list[tuple[str, str, bytes]] = []

    def execute(self, in_doc: InputDocument, raises_on_error: bool) -> ConversionResult:
        stream = in_doc._backend.path_or_stream
        self.calls.append((in_doc.format.value, in_doc.file.name, stream.getvalue()))
        result = ConversionResult(input=in_doc, status=self.status)
        result.document = DoclingDocument(name=in_doc.file.stem)
        result.document.add_text(
            label=DocItemLabel.TEXT,
            text=_TRANSCRIPT,
            source=TrackSource(start_time=0.0, end_time=1.0),
        )
        if self.status == ConversionStatus.FAILURE:
            result.errors.append(
                ErrorItem(
                    component_type=DoclingComponentType.PIPELINE,
                    module_name="StubMediaPipeline",
                    error_message="no audio track",
                )
            )
        return result


@pytest.fixture
def stub(monkeypatch: pytest.MonkeyPatch) -> _StubMediaPipeline:
    stub = _StubMediaPipeline()
    monkeypatch.setattr(SimplePipeline, "_get_media_pipeline", lambda self, kind: stub)
    return stub


def _deck(tmp_path: Path, file_name: str = "clip.mp4"):
    """Return a deck with a title and a video, and the video shape."""
    from pptx import Presentation
    from pptx.util import Inches

    media_path = tmp_path / file_name
    media_path.write_bytes(_MEDIA_BYTES)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
    title.text_frame.text = "Slide with media"
    video = slide.shapes.add_movie(
        str(media_path),
        Inches(1),
        Inches(2),
        Inches(4),
        Inches(2.25),
        mime_type="video/mp4",
    )
    return prs, slide, video


def _convert(
    path: Path,
    do_media_conversion: bool = True,
    backend_options: Optional[MsPowerpointBackendOptions] = None,
) -> ConversionResult:
    converter = DocumentConverter(
        allowed_formats=[InputFormat.PPTX],
        format_options={
            InputFormat.PPTX: PowerpointFormatOption(
                pipeline_options=ConvertPipelineOptions(
                    do_media_conversion=do_media_conversion
                ),
                backend_options=backend_options,
            )
        },
    )
    return converter.convert(path, raises_on_error=False)


def _media_group(doc: DoclingDocument, picture: PictureItem) -> Optional[GroupItem]:
    """Return the group right after the picture, if it holds converted media."""
    assert picture.parent is not None
    siblings = picture.parent.resolve(doc).children
    index = [ref.cref for ref in siblings].index(picture.self_ref)
    if index + 1 == len(siblings):
        return None
    sibling = siblings[index + 1].resolve(doc)
    if isinstance(sibling, GroupItem) and sibling.label == GroupLabel.SECTION:
        return sibling
    return None


def test_media_document_goes_after_its_picture(tmp_path: Path, stub):
    """The converted media follows the picture that shows it, on the same slide.

    Its items get the provenance of the picture: the media plays there.
    Without ``do_media_conversion`` nothing is converted.
    """
    prs, _, _ = _deck(tmp_path)
    deck = tmp_path / "deck.pptx"
    prs.save(deck)

    plain = _convert(deck, do_media_conversion=False).document
    assert stub.calls == []
    assert _media_group(plain, plain.pictures[0]) is None

    result = _convert(deck)
    assert result.status == ConversionStatus.SUCCESS
    assert stub.calls == [("video", "media1.mp4", _MEDIA_BYTES)]
    doc = result.document
    picture = doc.pictures[0]
    group = _media_group(doc, picture)
    assert group is not None
    assert group.name == "media: ppt/media/media1.mp4"
    transcript = group.children[0].resolve(doc)
    assert isinstance(transcript, TextItem)
    assert transcript.text == _TRANSCRIPT
    assert transcript.source[0].start_time == 0.0
    assert transcript.prov[0].page_no == 1
    assert transcript.prov[0].bbox == picture.prov[0].bbox
    assert transcript.prov[0].charspan == (0, len(_TRANSCRIPT))
    assert doc.export_to_markdown() == (
        "Slide with media\n\n<!-- image -->\n\nppt/media/media1.mp4\n\nSpoken words"
    )


def test_media_of_a_hidden_shape_stays_hidden(tmp_path: Path, stub):
    """The converted media takes the content layer of its picture."""
    prs, _, video = _deck(tmp_path)
    video._element.xpath("./*/p:cNvPr")[0].set("hidden", "1")
    deck = tmp_path / "hidden.pptx"
    prs.save(deck)

    doc = _convert(deck).document

    group = _media_group(doc, doc.pictures[0])
    assert group is not None
    assert group.content_layer == ContentLayer.INVISIBLE
    assert group.children[0].resolve(doc).content_layer == ContentLayer.INVISIBLE
    assert _TRANSCRIPT not in doc.export_to_markdown()


def test_linked_media_follows_the_fetch_options(tmp_path: Path, stub):
    """A linked file is loaded only when the backend options allow it.

    A file that is not allowed is skipped, and the conversion still succeeds.
    A local file must also be in the folder of the deck: PowerPoint links a
    local file by its absolute path, and such a link is refused and reported.
    """
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn

    def link_video(target: str, path: Path) -> None:
        prs, slide, video = _deck(tmp_path)
        r_id = slide.part.relate_to(target, RT.VIDEO, is_external=True)
        video._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile")[0].set(
            qn("r:link"), r_id
        )
        # Drop the p14:media extension, which points at the embedded copy.
        ext_lst = video._element.xpath("./p:nvPicPr/p:nvPr/p:extLst")[0]
        ext_lst.getparent().remove(ext_lst)
        prs.save(path)

    (tmp_path / "talk.mp4").write_bytes(b"linked video next to the deck")
    local_deck = tmp_path / "local.pptx"
    link_video("talk.mp4", local_deck)
    absolute_deck = tmp_path / "absolute.pptx"
    link_video((tmp_path / "talk.mp4").as_uri(), absolute_deck)
    remote_deck = tmp_path / "remote.pptx"
    link_video("https://example.com/talk.mp4", remote_deck)
    local_fetch = MsPowerpointBackendOptions(enable_local_fetch=True)

    for deck in (local_deck, absolute_deck, remote_deck):
        assert _convert(deck).status == ConversionStatus.SUCCESS
    assert stub.calls == []

    allowed = _convert(local_deck, backend_options=local_fetch)
    assert allowed.status == ConversionStatus.SUCCESS
    assert stub.calls == [("video", "talk.mp4", b"linked video next to the deck")]
    assert _media_group(allowed.document, allowed.document.pictures[0]) is not None

    absolute = _convert(absolute_deck, backend_options=local_fetch)
    assert absolute.status == ConversionStatus.PARTIAL_SUCCESS
    assert len(stub.calls) == 1


@pytest.mark.parametrize("missing_extra", [False, True])
def test_media_that_cannot_be_converted_is_reported(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, missing_extra: bool
):
    """A failed media conversion keeps the picture and makes a partial success.

    A missing ASR extra fails only the media too, the same way a missing ffmpeg
    fails the video pipeline.
    """
    failing = _StubMediaPipeline(status=ConversionStatus.FAILURE)

    def get_media_pipeline(self, kind):
        if missing_extra:
            raise ImportError("whisper is not installed")
        return failing

    monkeypatch.setattr(SimplePipeline, "_get_media_pipeline", get_media_pipeline)
    prs, _, _ = _deck(tmp_path)
    deck = tmp_path / "deck.pptx"
    prs.save(deck)

    result = _convert(deck)

    reason = "whisper is not installed" if missing_extra else "no audio track"
    assert result.status == ConversionStatus.PARTIAL_SUCCESS
    assert [error.error_message for error in result.errors] == [
        f"Media file ppt/media/media1.mp4 cannot be converted: {reason}"
    ]
    assert len(result.document.pictures) == 1
    assert _media_group(result.document, result.document.pictures[0]) is None


def test_odp_audio_is_converted(tmp_path: Path, stub):
    """The ODP backend loads its embedded media for the pipeline too."""
    pytest.importorskip("odfdo")
    from odfdo import Document as OdfDocument, DrawPage, Element, Frame

    from docling.document_converter import OdpFormatOption

    audio = tmp_path / "note.mp3"
    audio.write_bytes(_MEDIA_BYTES)
    odf = OdfDocument("presentation")
    media_path = odf.add_file(str(audio))
    frame = Frame(size=("2cm", "2cm"), position=("1cm", "1cm"))
    frame.append(
        Element.from_tag(
            f'<draw:plugin xlink:href="{media_path}" draw:mime-type="audio/mpeg"/>'
        )
    )
    odf.body.clear()
    page = DrawPage("page1", name="Slide One")
    page.append(frame)
    odf.body.append(page)
    path = tmp_path / "audio.odp"
    odf.save(str(path))

    converter = DocumentConverter(
        allowed_formats=[InputFormat.ODP],
        format_options={
            InputFormat.ODP: OdpFormatOption(
                pipeline_options=ConvertPipelineOptions(do_media_conversion=True)
            )
        },
    )
    doc = converter.convert(path).document

    assert stub.calls == [("audio", Path(media_path).name, _MEDIA_BYTES)]
    group = _media_group(doc, doc.pictures[0])
    assert group is not None
    assert group.children[0].resolve(doc).text == _TRANSCRIPT
