# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for PowerPoint chart plots: every plot type and combination charts."""

import datetime
import logging
from pathlib import Path
from typing import Any, Optional

import pytest
from docling_core.types.doc import PictureClassificationLabel, PictureItem

from docling.backend.mspowerpoint_backend import _serial_to_iso

from .test_backend_pptx import get_converter

_MONTHS = [datetime.date(2012, month, 1) for month in range(1, 13)]
_BARS = {
    "East": (85, 34, 100, 80, 116, 88, 71, 90, 121, 110, 101, 81),
    "West": (94, 72, 62, 91, 99, 70, 71, 46, 91, 57, 32, 157),
    "North": (91, 59, 85, 67, 130, 66, 51, 40, 62, 61, 61, 52),
    "South": (71, 110, 101, 28, 39, 52, 48, 125, 43, 122, 60, 63),
}
_LINES = {
    "Margin": tuple(
        n / 1000 for n in (126, 74, 110, 92, 86, 89, 77, 40, 101, 62, 81, 62)
    ),
    "Target": tuple(
        n / 1000 for n in (91, 75, 94, 72, 103, 75, 66, 81, 85, 94, 69, 95)
    ),
}


def _save_combo_chart_deck(
    path: Path, second_plot: str = "lineChart", linked_axis_title: bool = False
) -> None:
    """Save a deck whose chart draws columns and a second plot on a second axis.

    python-pptx writes single-plot charts only, so the chart starts as clustered
    columns over date categories. The last two series then move into a
    ``c:<second_plot>`` element on a secondary value axis, which is how
    PowerPoint stores a "Clustered Column - Line on Secondary Axis" chart. With
    ``second_plot="stockChart"`` the second plot is a stock plot, as in a
    "Volume-High-Low-Close" chart, for which python-pptx has no plot class.

    The second plot's category axis is hidden. It is a text axis with its own
    ``General`` number format, so it shows serial numbers, not dates. The
    category axis is titled "Region Position" and the value axes "Bar Sales".
    With ``linked_axis_title`` the category axis title is linked to a cell, so
    its text is in a cell cache instead of in rich text.
    """
    import copy

    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    chart_data = CategoryChartData()
    chart_data.categories = _MONTHS
    for name, values in {**_BARS, **_LINES}.items():
        chart_data.add_series(name, values)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5),
        Inches(0.5),
        Inches(9),
        Inches(6),
        chart_data,
    ).chart
    chart.category_axis.has_title = True
    chart.category_axis.axis_title.text_frame.text = "Region Position"
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "Bar Sales"

    plot_area = chart.element.chart.plotArea
    if linked_axis_title:
        title_text = plot_area.find(qn("c:dateAx")).find(qn("c:title")).find(qn("c:tx"))
        title_text.remove(title_text.find(qn("c:rich")))
        str_ref = etree.SubElement(title_text, qn("c:strRef"))
        etree.SubElement(str_ref, qn("c:f")).text = "Sheet1!$A$1"
        str_cache = etree.SubElement(str_ref, qn("c:strCache"))
        etree.SubElement(str_cache, qn("c:ptCount")).set("val", "1")
        point = etree.SubElement(str_cache, qn("c:pt"))
        point.set("idx", "0")
        etree.SubElement(point, qn("c:v")).text = "Region Position"
    bar_chart = plot_area.find(qn("c:barChart"))
    second = etree.Element(qn(f"c:{second_plot}"))
    if second_plot == "lineChart":
        etree.SubElement(second, qn("c:grouping")).set("val", "standard")
        etree.SubElement(second, qn("c:varyColors")).set("val", "0")
    for ser in bar_chart.findall(qn("c:ser"))[len(_BARS) :]:
        for invert in ser.findall(qn("c:invertIfNegative")):
            ser.remove(invert)
        second.append(ser)
    if second_plot == "lineChart":
        etree.SubElement(second, qn("c:marker")).set("val", "1")
    else:
        etree.SubElement(second, qn("c:hiLowLines"))
    etree.SubElement(second, qn("c:axId")).set("val", "900")
    etree.SubElement(second, qn("c:axId")).set("val", "901")
    bar_chart.addnext(second)

    value_axis = plot_area.find(qn("c:valAx"))
    secondary_category_axis = etree.Element(qn("c:catAx"))
    etree.SubElement(secondary_category_axis, qn("c:axId")).set("val", "900")
    scaling = etree.SubElement(secondary_category_axis, qn("c:scaling"))
    etree.SubElement(scaling, qn("c:orientation")).set("val", "minMax")
    etree.SubElement(secondary_category_axis, qn("c:delete")).set("val", "1")
    etree.SubElement(secondary_category_axis, qn("c:axPos")).set("val", "b")
    number_format = etree.SubElement(secondary_category_axis, qn("c:numFmt"))
    number_format.set("formatCode", "General")
    number_format.set("sourceLinked", "0")
    etree.SubElement(secondary_category_axis, qn("c:crossAx")).set("val", "901")
    secondary_value_axis = copy.deepcopy(value_axis)
    secondary_value_axis.find(qn("c:axId")).set("val", "901")
    secondary_value_axis.find(qn("c:crossAx")).set("val", "900")
    value_axis.addnext(secondary_value_axis)
    secondary_value_axis.addnext(secondary_category_axis)
    prs.save(path)


def _single_plot_deck(plot: str) -> tuple[Any, Any]:
    """Return a deck whose chart has one ``c:<plot>`` plot, and that plot element.

    The plot has one series "2026" with 10, 20 and 30 over Q1 to Q3. python-pptx
    writes only the plot types that it has a class for, so the series and axis
    ids of a clustered column chart move into a new ``c:<plot>`` element.
    """
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3"]
    chart_data.add_series("2026", (10, 20, 30))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    ).chart
    bar_chart = chart.element.chart.plotArea.find(qn("c:barChart"))
    other_plot = etree.Element(qn(f"c:{plot}"))
    other_plot.extend(bar_chart.findall(qn("c:ser")) + bar_chart.findall(qn("c:axId")))
    bar_chart.addnext(other_plot)
    bar_chart.getparent().remove(bar_chart)
    return prs, other_plot


def _only_chart(path: Path) -> PictureItem:
    doc = get_converter().convert(path).document
    pictures = list(doc.pictures)
    assert len(pictures) == 1
    return pictures[0]


def _grid(picture: PictureItem) -> list[list[str]]:
    assert picture.meta is not None and picture.meta.tabular_chart is not None
    return [
        [cell.text for cell in row]
        for row in picture.meta.tabular_chart.chart_data.grid
    ]


def _labels(picture: PictureItem) -> list[str]:
    assert picture.meta is not None and picture.meta.classification is not None
    return [p.class_name for p in picture.meta.classification.predictions]


def test_combo_chart_lists_each_plot_type_and_keeps_all_series(tmp_path: Path):
    """A column + line chart is classified as both, back-most plot first.

    The main prediction stays the column plot, as before, so serializers that
    read only the main prediction still treat the picture as a bar chart. The
    line series on the secondary axis stay in the data table, after the column
    series.
    """
    path = tmp_path / "combo_chart.pptx"
    _save_combo_chart_deck(path)

    picture = _only_chart(path)

    assert _labels(picture) == [
        PictureClassificationLabel.BAR_CHART.value,
        PictureClassificationLabel.LINE_CHART.value,
    ]
    grid = _grid(picture)
    assert grid[0][1:] == ["East", "West", "North", "South", "Margin", "Target"]
    assert grid[1][1:] == ["85", "94", "91", "71", "0.126", "0.091"]
    assert len(grid) == 1 + len(_MONTHS)


def test_stock_plot_keeps_its_series_next_to_the_other_plots(tmp_path: Path):
    """The series of a stock plot follow the column series in the data table.

    python-pptx has no plot class for a stock plot. Before, its series were
    skipped with a warning. Now they are read, and the stock plot counts as
    another chart type.
    """
    path = tmp_path / "stock_combo_chart.pptx"
    _save_combo_chart_deck(path, second_plot="stockChart")

    picture = _only_chart(path)

    assert _labels(picture) == [
        PictureClassificationLabel.BAR_CHART.value,
        PictureClassificationLabel.OTHER_CHART.value,
    ]
    grid = _grid(picture)
    assert grid[0][1:] == [*_BARS, *_LINES]
    assert grid[1][1:] == ["85", "94", "91", "71", "0.126", "0.091"]


@pytest.mark.parametrize(
    ("plot", "label"),
    [
        ("bar3DChart", PictureClassificationLabel.BAR_CHART),
        ("line3DChart", PictureClassificationLabel.LINE_CHART),
        ("pie3DChart", PictureClassificationLabel.PIE_CHART),
        ("area3DChart", PictureClassificationLabel.OTHER_CHART),
        ("ofPieChart", PictureClassificationLabel.OTHER_CHART),
        ("stockChart", PictureClassificationLabel.OTHER_CHART),
        ("surfaceChart", PictureClassificationLabel.OTHER_CHART),
        ("surface3DChart", PictureClassificationLabel.OTHER_CHART),
    ],
)
def test_plot_without_a_python_pptx_class_keeps_its_data(
    tmp_path: Path, plot: str, label: PictureClassificationLabel
):
    """A plot that python-pptx has no class for is classified by its tag.

    Before, such a chart was "other_chart" with no data table, and a 3-D area
    chart stopped the conversion of the whole presentation, because python-pptx
    has a plot class but no series class for it. Now the chart keeps its type
    and its data.
    """
    prs, _ = _single_plot_deck(plot)
    path = tmp_path / f"{plot}.pptx"
    prs.save(path)

    picture = _only_chart(path)

    assert _labels(picture) == [label.value]
    assert _grid(picture) == [["", "2026"], ["Q1", "10"], ["Q2", "20"], ["Q3", "30"]]


def test_malformed_series_drops_only_the_chart_data(tmp_path: Path, caplog):
    """A series without ``c:order`` loses the chart's data, not the presentation.

    python-pptx raises ``InvalidXmlError`` when a required element is missing.
    That error must not stop the conversion of the whole presentation. The
    chart keeps its type without a data table, and a warning names the missing
    element.
    """
    from pptx.oxml.ns import qn

    prs, stock_plot = _single_plot_deck("stockChart")
    series = stock_plot.find(qn("c:ser"))
    series.remove(series.find(qn("c:order")))
    path = tmp_path / "stock_chart_without_order.pptx"
    prs.save(path)

    with caplog.at_level(
        logging.WARNING, logger="docling.backend.mspowerpoint_backend"
    ):
        picture = _only_chart(path)

    assert _labels(picture) == [PictureClassificationLabel.OTHER_CHART.value]
    assert picture.meta is not None and picture.meta.tabular_chart is None
    assert "c:order" in caplog.text


def test_date_axis_categories_are_iso_dates(tmp_path: Path):
    """Categories on a date axis read as ISO 8601 dates, not as serial numbers.

    The chart caches 1 January 2012 as the serial 40909. The column plot is
    drawn against a date axis, while the hidden category axis of the line plot
    is a text axis. The dates therefore depend on finding the column plot's own
    axis.
    """
    path = tmp_path / "combo_chart.pptx"
    _save_combo_chart_deck(path)

    grid = _grid(_only_chart(path))

    assert [row[0] for row in grid[1:]] == [month.isoformat() for month in _MONTHS]


@pytest.mark.parametrize("linked", [False, True], ids=["rich-text", "cell-linked"])
def test_category_axis_title_heads_the_category_column(tmp_path: Path, linked: bool):
    """The category axis title fills the top-left cell of the data table.

    That cell was always empty. A title linked to a cell is read from the cell
    cache. The value axis titles have no column to head, so they stay out.
    """
    path = tmp_path / "combo_chart.pptx"
    _save_combo_chart_deck(path, linked_axis_title=linked)

    grid = _grid(_only_chart(path))

    assert grid[0] == ["Region Position", *_BARS, *_LINES]


@pytest.mark.parametrize(
    ("axis_format", "source_linked", "cache_format", "expected"),
    [
        ("m/d/yy", "0", "General", ["2012-01-01", "2012-02-01"]),
        ("General", "1", "mmm d, yy", ["2012-01-01", "2012-02-01"]),
        (None, None, "mmm d, yy", ["2012-01-01", "2012-02-01"]),
        ("General", "0", "mmm d, yy", ["40909", "40940"]),
        ("h:mm", "0", "General", ["40909", "40940"]),
    ],
    ids=["axis-date", "linked-date", "no-axis-format", "axis-general", "axis-time"],
)
def test_text_axis_categories_follow_the_shown_number_format(
    tmp_path: Path,
    axis_format: Optional[str],
    source_linked: Optional[str],
    cache_format: str,
    expected: list[str],
):
    """Serial categories on a text axis become dates only when shown as dates.

    The axis shows its own number format, unless it links the format to the
    source data or has none. A time-only format does not make a date.
    """
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    chart_data = CategoryChartData()
    chart_data.categories = [40909, 40940]
    chart_data.categories.number_format = cache_format
    chart_data.add_series("Sales", (10, 20))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    ).chart
    if axis_format is not None and source_linked is not None:
        number_format = etree.Element(qn("c:numFmt"))
        number_format.set("formatCode", axis_format)
        number_format.set("sourceLinked", source_linked)
        category_axis = chart.element.chart.plotArea.find(qn("c:catAx"))
        category_axis.find(qn("c:axPos")).addnext(number_format)
    path = tmp_path / "serial_categories.pptx"
    prs.save(path)

    grid = _grid(_only_chart(path))

    assert [row[0] for row in grid[1:]] == expected


@pytest.mark.parametrize(
    ("text", "date1904", "expected"),
    [
        ("40909", False, "2012-01-01"),
        ("40909.0", False, "2012-01-01"),
        ("39447", True, "2012-01-01"),
        ("40909.5", False, "2012-01-01T12:00:00"),
        ("1", False, "1900-01-01"),
        ("61", False, "1900-03-01"),
        ("Jan", False, "Jan"),
        ("", False, ""),
        ("-1", False, "-1"),
        ("1e300", False, "1e300"),
    ],
)
def test_serial_to_iso(text: str, date1904: bool, expected: str):
    """Excel serials map to ISO dates in both date systems; other text stays."""
    assert _serial_to_iso(text, date1904) == expected
