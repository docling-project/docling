# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Check the DOCX heading tree; core tests cover its exports and chunking."""

from pathlib import Path

import pytest
from docling_core.types.doc import (
    DoclingDocument,
    FormulaItem,
    InlineGroup,
    SectionHeaderItem,
    TextItem,
)
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from docling.backend.msword_backend import MsWordDocumentBackend
from docling.datamodel.base_models import InputFormat
from docling.datamodel.document import InputDocument


@pytest.mark.parametrize("level", [1, 3])
@pytest.mark.parametrize(
    ("before", "after"),
    [("Heading eq ", ""), ("", " after"), ("Before ", " after"), ("", "")],
)
def test_heading_equation_has_inline_content(
    tmp_path: Path, level: int, before: str, after: str
) -> None:
    document = Document()
    heading = document.add_heading(before, level)
    heading._p.append(
        parse_xml(f"<m:oMath {nsdecls('m')}><m:r><m:t>y=2</m:t></m:r></m:oMath>")
    )
    heading.add_run(after)
    document.add_paragraph("Body under it.")
    document.add_heading("Next heading", level)
    document.add_paragraph("Next body.")
    path = tmp_path / "heading-equation.docx"
    document.save(path)
    doc = InputDocument(
        path_or_stream=path, format=InputFormat.DOCX, backend=MsWordDocumentBackend
    )._backend.convert()
    doc = DoclingDocument.model_validate(doc.export_to_dict())

    headers = [item for item in doc.texts if isinstance(item, SectionHeaderItem)]
    header, next_header = headers
    assert header.text == header.orig == ""
    assert header.level == level
    inline = header.children[0].resolve(doc)
    assert isinstance(inline, InlineGroup)
    assert inline.parent == header.get_ref()
    content = [ref.resolve(doc) for ref in inline.children]
    assert all(isinstance(item, TextItem) for item in content)
    assert [item.text.strip() for item in content] == (
        ([before.strip()] if before else [])
        + ["y=2"]
        + ([after.strip()] if after else [])
    )
    formulas = [item for item in content if isinstance(item, FormulaItem)]
    assert len(formulas) == 1
    assert formulas[0].text == formulas[0].orig == "y=2"
    assert all(item.parent == inline.get_ref() for item in content)
    body = next(item for item in doc.texts if item.text == "Body under it.")
    assert body.parent == header.get_ref()
    assert next_header.text == "Next heading"
    next_body = next(item for item in doc.texts if item.text == "Next body.")
    assert next_body.parent == next_header.get_ref()
    assert all(
        "<eq>" not in item.text and "</eq>" not in item.text for item in doc.texts
    )


def test_heading_keeps_repeated_equations_in_order(tmp_path: Path) -> None:
    document = Document()
    heading = document.add_heading("Start ", 1)
    heading.add_run("secret").font.hidden = True
    math = f"<m:oMath {nsdecls('m')}><m:r><m:t>x=1</m:t></m:r></m:oMath>"
    heading._p.append(parse_xml(math))
    heading.add_run(" middle ")
    heading._p.append(parse_xml(math))
    heading.add_run(" end")
    path = tmp_path / "heading-equations.docx"
    document.save(path)
    doc = InputDocument(
        path_or_stream=path, format=InputFormat.DOCX, backend=MsWordDocumentBackend
    )._backend.convert()

    header = next(item for item in doc.texts if isinstance(item, SectionHeaderItem))
    assert header.text == ""
    inline = header.children[0].resolve(doc)
    assert isinstance(inline, InlineGroup)
    content = [ref.resolve(doc) for ref in inline.children]
    assert [item.text.strip() for item in content] == [
        "Start",
        "x=1",
        "middle",
        "x=1",
        "end",
    ]
    assert [isinstance(item, FormulaItem) for item in content] == [
        False,
        True,
        False,
        True,
        False,
    ]
    assert all("secret" not in item.text for item in doc.texts)


def test_numbered_heading_equation_keeps_the_numbering(tmp_path: Path) -> None:
    document = Document()
    num_id = document.styles["List Number"].element.pPr.numPr.numId.val
    for title, with_equation in [("First", False), ("Value ", True), ("Last", False)]:
        heading = document.add_heading(title, 1)
        heading._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_id
        if with_equation:
            heading._p.append(
                parse_xml(
                    f"<m:oMath {nsdecls('m')}><m:r><m:t>y=2</m:t></m:r></m:oMath>"
                )
            )
    path = tmp_path / "numbered-heading-equation.docx"
    document.save(path)
    doc = InputDocument(
        path_or_stream=path, format=InputFormat.DOCX, backend=MsWordDocumentBackend
    )._backend.convert()

    headers = [item for item in doc.texts if isinstance(item, SectionHeaderItem)]
    assert [header.text for header in headers] == ["1 First", "", "3 Last"]
    inline = headers[1].children[0].resolve(doc)
    assert isinstance(inline, InlineGroup)
    assert [ref.resolve(doc).text.strip() for ref in inline.children] == [
        "2 Value",
        "y=2",
    ]
